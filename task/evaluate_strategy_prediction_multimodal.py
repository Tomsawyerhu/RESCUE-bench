import argparse
import json
import os
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
import requests
try:
    from tqdm.auto import tqdm
except Exception:  # pragma: no cover - tqdm is optional for script progress.
    def tqdm(iterable=None, **kwargs):
        return iterable if iterable is not None else _NullProgress()


    class _NullProgress:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def update(self, n=1):
            return None

        def set_postfix_str(self, *args, **kwargs):
            return None


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from task.evaluate_strategy_prediction import (
    DEFAULT_API_BASE,
    STRATEGY_DEFINITIONS_BY_MODE,
    batched,
    call_json_model,
    canonical_target,
    clean_text,
    current_row_strategy,
    ensure_unique_ordered,
    group_label_for_sample,
    iter_prediction_items,
    load_data_samples,
    load_sample_id_filter,
    mask_sensitive_terms,
    normalize_api_base,
    rank_metrics,
    strategy_definitions_for_sample,
    target_label,
)
from task.multimodal_eval_common import (
    DEFAULT_UPLOAD_CACHE,
    MODE_ALIASES,
    VALID_MODES,
    build_prior_context_rows,
    call_video_json_model_with_sanitize_retry,
    ensure_upload_url,
    load_review_manifest,
    prompt_record_for_mode,
    row_timestamp,
)
from utils.project_env import load_project_env


load_project_env(ROOT_DIR)


TASK_NAME = "predict_next_therapist_strategy_from_prior_context_with_known_support_target"


def normalize_internal_emotion_items(raw_items: Sequence[Dict[str, Any]], fallback_person: str = "") -> List[Dict[str, Any]]:
    normalized: List[Dict[str, Any]] = []
    for item in raw_items or []:
        if not isinstance(item, dict):
            continue
        emotion_description = clean_text(
            item.get("emotion_description") or item.get("emotion") or item.get("description")
        )
        if not emotion_description:
            continue
        payload = {
            "person": clean_text(item.get("person")) or fallback_person,
            "emotion_description": emotion_description,
        }
        try:
            intensity = int(round(float(item.get("intensity"))))
        except Exception:
            intensity = None
        if intensity is not None:
            payload["intensity"] = intensity
        normalized.append(payload)
    return normalized


def find_recent_target_internal_emotion(
    rows: Sequence[Dict[str, Any]],
    before_row_index: int,
    context_size: int,
    target_label_text: str,
    participants: Sequence[str],
    group_label: str,
) -> List[Dict[str, Any]]:
    if before_row_index <= 1 or context_size == 0:
        return []

    context_start = 1 if context_size < 0 else max(1, before_row_index - context_size)
    participant_set = {clean_text(item) for item in participants if clean_text(item)}

    for prior_index in range(before_row_index - 1, context_start - 1, -1):
        row = rows[prior_index - 1]
        items = normalize_internal_emotion_items(
            row.get("internal_emotion") or [],
            fallback_person=clean_text(row.get("primary_speaker")),
        )
        if not items:
            continue
        if target_label_text == group_label:
            selected = [item for item in items if clean_text(item.get("person")) in participant_set]
            if selected:
                return selected
            continue
        selected = [item for item in items if clean_text(item.get("person")) == target_label_text]
        if selected:
            return selected
    return []


def build_targets(
    sample: Dict[str, Any],
    context_size: int,
    require_video: bool,
) -> List[Dict[str, Any]]:
    rows = sample["rows"]
    manifest_rows = load_review_manifest(sample["sample_id"])
    targets: List[Dict[str, Any]] = []
    strategy_definitions = strategy_definitions_for_sample(sample)
    group_label = group_label_for_sample(sample)

    for row_index, row in enumerate(rows, start=1):
        if clean_text(row.get("primary_speaker")) != "Therapist":
            continue

        gold_strategy, gold_target, gold_target_source = current_row_strategy(row, sample["group_mode"])
        if not gold_strategy or gold_strategy not in strategy_definitions:
            continue

        if not gold_target:
            fallback_target = canonical_target(row.get("target") or [])
            if fallback_target:
                gold_target = fallback_target
                gold_target_source = "row_target_fallback"
            else:
                continue

        gold_target_label = target_label(gold_target, group_label=group_label)
        if not gold_target_label:
            continue

        prior_context = build_prior_context_rows(rows, manifest_rows, row_index, context_size)
        if require_video and not any(clean_text(item.get("clip_path")) for item in prior_context):
            continue

        targets.append(
            {
                "sample_id": sample["sample_id"],
                "couple": sample["couple"],
                "file_rel": sample["file_rel"],
                "turn_id": "%s__row%d" % (Path(sample["file_rel"]).stem, row_index),
                "row_index": row_index,
                "timestamp": row_timestamp(row),
                "dialogue": clean_text(row.get("dialogue_cleaned")),
                "gold_strategy": gold_strategy,
                "gold_target_label": gold_target_label,
                "gold_target_source": gold_target_source,
                "participants": list(sample["participants"]),
                "prior_context": prior_context,
                "known_target_recent_internal_emotion": find_recent_target_internal_emotion(
                    rows=rows,
                    before_row_index=row_index,
                    context_size=context_size,
                    target_label_text=gold_target_label,
                    participants=sample["participants"],
                    group_label=group_label,
                ),
                "group_mode": sample["group_mode"],
            }
        )

    return targets


def strategy_prompt_record(
    target: Dict[str, Any],
    mode: str,
    context_size: int,
    include_target_internal_emotion: bool,
) -> Dict[str, Any]:
    record: Dict[str, Any] = {
        "turn_id": target["turn_id"],
        "row_index": target["row_index"],
        "participants": target["participants"],
        "known_support_target": target["gold_target_label"],
    }
    if context_size != 0:
        record["prior_context"] = [prompt_record_for_mode(item, mode) for item in target["prior_context"]]
    else:
        record["prior_context"] = []
    if include_target_internal_emotion and target.get("known_target_recent_internal_emotion"):
        record["known_support_target_internal_emotion"] = target["known_target_recent_internal_emotion"]
    return record


def sanitize_strategy_records(records: Sequence[Dict[str, Any]], level: str) -> List[Dict[str, Any]]:
    sanitized: List[Dict[str, Any]] = []
    for record in records:
        item = json.loads(json.dumps(record, ensure_ascii=False))
        if level == "none":
            sanitized.append(item)
            continue
        for ctx in item.get("prior_context", []) or []:
            if "dialogue_cleaned" in ctx:
                if level == "mask":
                    ctx["dialogue_cleaned"] = mask_sensitive_terms(clean_text(ctx.get("dialogue_cleaned")))
                else:
                    ctx["dialogue_cleaned"] = "[redacted sensitive prior dialogue]"
            for bg in ctx.get("background_dialogue", []) or []:
                if "content" in bg:
                    if level == "mask":
                        bg["content"] = mask_sensitive_terms(clean_text(bg.get("content")))
                    else:
                        bg["content"] = "[redacted sensitive background dialogue]"
        sanitized.append(item)
    return sanitized


def build_turn_video_prompt_bundle(
    target: Dict[str, Any],
    context_size: int,
    include_target_internal_emotion: bool,
    upload_cache: Dict[str, Any],
    upload_cache_path: Path,
    api_key: str,
    model: str,
) -> tuple[Dict[str, Any], List[str]]:
    record = strategy_prompt_record(target, "turn_video_audio", context_size, include_target_internal_emotion)
    video_urls: List[str] = []
    attachment_index = 1
    enriched_prior_context: List[Dict[str, Any]] = []

    for source_item, prompt_item in zip(target["prior_context"], record.get("prior_context", [])):
        enriched_item = dict(prompt_item)
        clip_path = clean_text(source_item.get("clip_path"))
        if clip_path and Path(clip_path).exists():
            video_urls.append(ensure_upload_url(upload_cache, upload_cache_path, api_key, model, Path(clip_path)))
            enriched_item["video_attachment_index"] = attachment_index
            try:
                duration = float(source_item.get("clip_duration_sec") or 0.0)
            except Exception:
                duration = 0.0
            if duration > 0:
                enriched_item["video_duration_sec"] = round(duration, 3)
            attachment_index += 1
        enriched_prior_context.append(enriched_item)

    record["prior_context"] = enriched_prior_context
    return record, video_urls


def prediction_prompt(
    sample_id: str,
    group_mode: str,
    mode: str,
    records: Sequence[Dict[str, Any]],
    context_size: int,
    include_target_internal_emotion: bool,
) -> str:
    strategy_definitions = STRATEGY_DEFINITIONS_BY_MODE[group_mode]
    group_label = "Family" if group_mode == "family" else "Couple"
    session_phrase = "family-therapy interaction" if group_mode == "family" else "couples-therapy interaction"
    mode_guidance = {
        "dialogue": "Use only language-side evidence from prior_context: primary_speaker, target, dialogue_cleaned, and background_dialogue.",
        "audio": "Use prior_context language-side evidence together with tone_of_voice. Do not use visual annotations or video evidence.",
        "visual": "Use prior_context language plus visual annotation evidence from body_posture, facial_expressions, self_directed_behavior, and interaction_behavior.",
        "audiovisual": "Use prior_context language, visual annotations, and tone_of_voice.",
        "turn_video_audio": "Use the attached row-level prior-context video clips together with the prior_context transcript records. The clips already contain audio, so do not expect a separate tone_of_voice field.",
    }
    context_guidance = (
        "No prior dialogue context is provided.\n"
        if context_size == 0
        else (
            "Use all available previous rows in prior_context. Do not use any future information from the unseen therapist turn.\n"
            if context_size < 0
            else f"Use at most the previous {context_size} rows in prior_context. Do not use any future information from the unseen therapist turn.\n"
        )
    )
    video_guidance = (
        ""
        if mode != "turn_video_audio"
        else "For turn_video_audio, attached video clips appear in attachment order. Any prior_context item with video_attachment_index refers to its matching attached clip, ordered from oldest to newest.\n"
    )
    single_target_guidance = (
        ""
        if mode != "turn_video_audio"
        else "For turn_video_audio, the input contains exactly one target. You must return exactly one prediction item for that target.\n"
    )
    target_internal_emotion_guidance = (
        "If known_support_target_internal_emotion is provided, treat it as a gold annotation of the target user's recent internal emotional state and use it as a strong cue.\n"
        if include_target_internal_emotion
        else ""
    )
    return (
        f"You are predicting the therapist's NEXT supportive strategy in a {session_phrase}.\n"
        + "You only see PRIOR CONTEXT before the target therapist turn. You must not assume access to the upcoming therapist utterance.\n"
        + "The support target is GIVEN to you as known_support_target. Use it as a conditioning signal.\n"
        + target_internal_emotion_guidance
        + "Your task is only to rank the 3 most likely therapist strategies from this closed set:\n"
        + f"{json.dumps(strategy_definitions, ensure_ascii=False)}\n"
        + context_guidance
        + video_guidance
        + single_target_guidance
        + f"Mode: {mode}\n"
        + f"{mode_guidance[mode]}\n"
        + "Important rules:\n"
        + "- Predict the MAIN FUNCTION of the likely next therapist move, not every plausible function.\n"
        + "- Prefer the most specific justified label. `track` is the broadest active label and should win only when no more specific label clearly fits.\n"
        + "- Because known_support_target is given, use it to disambiguate between overlapping strategies.\n"
        + (
            "- `goal_align` means reorienting toward shared task or purpose; generic process summary is not enough.\n"
            "- `detach` means helping participants step back together and view the problem as a shared pattern.\n"
            "- `join` means fostering softer emotional contact or mutual receptivity.\n"
            "- `repair` means reconnecting after strain, rupture, apology, misunderstanding, or disconnection.\n"
            "- `counterbalance` restores participation or alliance balance; `safeguard` protects a vulnerable person from being overrun.\n"
            "- `reframe` changes the meaning of the problem into a relational frame; simple sequence description is not enough.\n"
            "- `enact` means moving one participant toward direct in-room communication with another participant.\n"
            if group_mode == "couple"
            else "- `join` means alliance-building and keeping the family engaged enough to keep working.\n"
            "- `track` means explicitly following the live family pattern, coalition, or recurrent sequence.\n"
            "- `counterbalance` restores participation or influence to a less-heard family member.\n"
            "- `boundary` means clarifying or restructuring boundaries, roles, hierarchy, coalitions, or subsystem functioning.\n"
            "- `reframe` means shifting blame into a systemic family frame or shared challenge.\n"
            "- `repair` means apology, acknowledgement of hurt, reconnection, or healing after rupture.\n"
            "- `enact` means moving family members into direct in-room communication with each other.\n"
        )
        + "- Return exactly 3 distinct strategy labels.\n"
        + "- Do not output empty, none, or labels outside the closed set.\n"
        + "- Echo the exact Sample value and exact input turn_id values from Targets. Never copy placeholder or example IDs.\n"
        + "Output JSON only with this schema:\n"
        + "{"
        + "\"sample_id\":\"<echo the exact Sample value below>\","
        + "\"predictions\":["
        + "{"
        + "\"turn_id\":\"<echo one exact turn_id from Targets>\","
        + "\"strategy_top3\":[\"track\",\"evoke\",\"reframe\"],"
        + "\"reason\":\"short English reason based only on the prior context and known support target\""
        + "}"
        + "]"
        + "}\n"
        + "No markdown. JSON only.\n"
        + f"Sample: {sample_id}\n"
        + f"Targets: {json.dumps(list(records), ensure_ascii=False)}"
    )


def request_text_predictions(
    *,
    api_base: str,
    api_key: str,
    model: str,
    sample_id: str,
    group_mode: str,
    targets: Sequence[Dict[str, Any]],
    mode: str,
    context_size: int,
    include_target_internal_emotion: bool,
    max_tokens: int,
    disable_thinking: bool,
) -> Dict[str, Any]:
    prompts = [
        prediction_prompt(
            sample_id, group_mode, mode, sanitize_strategy_records(targets, "none"), context_size, include_target_internal_emotion
        ),
        prediction_prompt(
            sample_id, group_mode, mode, sanitize_strategy_records(targets, "mask"), context_size, include_target_internal_emotion
        ),
        prediction_prompt(
            sample_id, group_mode, mode, sanitize_strategy_records(targets, "redact"), context_size, include_target_internal_emotion
        ),
    ]
    last_error: Optional[Exception] = None
    for prompt_index, prompt in enumerate(prompts):
        for attempt in range(1, 4):
            try:
                return call_json_model(
                    api_base=api_base,
                    api_key=api_key,
                    model=model,
                    prompt=prompt,
                    max_tokens=max_tokens,
                    disable_thinking=disable_thinking,
                )
            except requests.HTTPError as exc:
                last_error = exc
                if "data_inspection_failed" in str(exc) and prompt_index < len(prompts) - 1:
                    break
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
            except (requests.RequestException, OSError) as exc:
                last_error = exc
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
    if last_error is not None:
        raise last_error
    raise RuntimeError("Unknown error while requesting text predictions")


def normalize_strategy_top3(values: Sequence[Any], strategy_definitions: Dict[str, str]) -> List[str]:
    labels = [label for label in ensure_unique_ordered(values) if label in strategy_definitions]
    while len(labels) < 3:
        labels.append("")
    return labels[:3]


def extract_prediction_items(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    items = iter_prediction_items(payload)
    if not items and isinstance(payload, dict) and payload.get("strategy_top3"):
        items = [payload]
    return items


def coerce_single_target_prediction_items(
    payload: Dict[str, Any],
    expected_turn_id: str,
    strategy_definitions: Dict[str, str],
) -> List[Dict[str, Any]]:
    items = extract_prediction_items(payload)
    if len(items) != 1:
        return items
    item = dict(items[0])
    if not normalize_strategy_top3(item.get("strategy_top3", []), strategy_definitions)[0]:
        return items
    item["turn_id"] = expected_turn_id
    return [item]


def request_turn_video_prediction(
    *,
    api_base: str,
    api_key: str,
    model: str,
    target: Dict[str, Any],
    include_target_internal_emotion: bool,
    upload_cache: Dict[str, Any],
    upload_cache_path: Path,
    context_size: int,
    max_tokens: int,
    disable_thinking: bool,
) -> Dict[str, Any]:
    prompt_record, video_urls = build_turn_video_prompt_bundle(
        target=target,
        context_size=context_size,
        include_target_internal_emotion=include_target_internal_emotion,
        upload_cache=upload_cache,
        upload_cache_path=upload_cache_path,
        api_key=api_key,
        model=model,
    )
    if not video_urls:
        raise RuntimeError(f"No prior-context video clips available for {target['turn_id']}")

    payload = call_video_json_model_with_sanitize_retry(
        api_base=api_base,
        api_key=api_key,
        model=model,
        video_urls=video_urls,
        prompt=prediction_prompt(
            target["sample_id"],
            clean_text(target.get("group_mode") or "couple"),
            "turn_video_audio",
            [prompt_record],
            context_size,
            include_target_internal_emotion,
        ),
        max_tokens=max_tokens,
        disable_thinking=disable_thinking,
    )
    raw_items = iter_prediction_items(payload)
    if not raw_items and isinstance(payload, dict) and payload.get("strategy_top3"):
        raw_items = [payload]
    for item in raw_items:
        turn_id = clean_text(item.get("turn_id")) or target["turn_id"]
        if turn_id == target["turn_id"]:
            return item
    raise RuntimeError(f"Missing video strategy prediction for {target['turn_id']}")


def summarize_rows(rows: Sequence[Dict[str, Any]]) -> Dict[str, float]:
    count = len(rows)
    if count == 0:
        return {
            "turn_count": 0,
            "strategy_recall_at_1": 0.0,
            "strategy_recall_at_2": 0.0,
            "strategy_recall_at_3": 0.0,
            "strategy_mrr": 0.0,
        }
    return {
        "turn_count": count,
        "strategy_recall_at_1": sum(row["strategy_recall_at_1"] for row in rows) / count,
        "strategy_recall_at_2": sum(row["strategy_recall_at_2"] for row in rows) / count,
        "strategy_recall_at_3": sum(row["strategy_recall_at_3"] for row in rows) / count,
        "strategy_mrr": sum(row["strategy_mrr"] for row in rows) / count,
    }


def build_sample_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_sample: Dict[tuple[str, str], List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_sample[(row["couple"], row["sample_id"])].append(row)
    return [
        {"couple": couple, "sample_id": sample_id, **summarize_rows(rows)}
        for (couple, sample_id), rows in sorted(by_sample.items(), key=lambda item: item[0][1])
    ]


def build_couple_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_couple: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_couple[row["couple"]].append(row)
    return [{"scope": couple, **summarize_rows(rows)} for couple, rows in sorted(by_couple.items(), key=lambda item: item[0])]


def build_strategy_confusion_rows(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    counter: Counter = Counter()
    for row in detail_rows:
        counter[(row["gold_strategy"], row["pred_strategy_top1"])] += 1
    return [
        {"gold_strategy": gold, "pred_strategy_top1": pred, "count": count}
        for (gold, pred), count in counter.most_common()
    ]


def write_result_workbook(
    output_path: Path,
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    detail_rows: Sequence[Dict[str, Any]],
    strategy_confusion_rows: Sequence[Dict[str, Any]],
) -> None:
    workbook = Workbook()

    ws_overall = workbook.active
    ws_overall.title = "overall"
    ws_overall.append(
        [
            "scope",
            "turn_count",
            "strategy_recall_at_1",
            "strategy_recall_at_2",
            "strategy_recall_at_3",
            "strategy_mrr",
        ]
    )
    for row in overall_rows:
        ws_overall.append(
            [
                row["scope"],
                row["turn_count"],
                row["strategy_recall_at_1"],
                row["strategy_recall_at_2"],
                row["strategy_recall_at_3"],
                row["strategy_mrr"],
            ]
        )

    ws_couples = workbook.create_sheet("couples")
    ws_couples.append(list(ws_overall.iter_rows(min_row=1, max_row=1, values_only=True))[0])
    for row in couple_rows:
        ws_couples.append(
            [
                row["scope"],
                row["turn_count"],
                row["strategy_recall_at_1"],
                row["strategy_recall_at_2"],
                row["strategy_recall_at_3"],
                row["strategy_mrr"],
            ]
        )

    ws_samples = workbook.create_sheet("samples")
    ws_samples.append(
        [
            "couple",
            "sample_id",
            "turn_count",
            "strategy_recall_at_1",
            "strategy_recall_at_2",
            "strategy_recall_at_3",
            "strategy_mrr",
        ]
    )
    for row in sample_rows:
        ws_samples.append(
            [
                row["couple"],
                row["sample_id"],
                row["turn_count"],
                row["strategy_recall_at_1"],
                row["strategy_recall_at_2"],
                row["strategy_recall_at_3"],
                row["strategy_mrr"],
            ]
        )

    ws_details = workbook.create_sheet("details")
    ws_details.append(
        [
            "couple",
            "sample_id",
            "turn_id",
            "row_index",
            "timestamp",
            "known_support_target",
            "gold_target_source",
            "gold_strategy",
            "pred_strategy_top1",
            "pred_strategy_top2",
            "pred_strategy_top3",
            "strategy_gold_rank",
            "strategy_recall_at_1",
            "strategy_recall_at_2",
            "strategy_recall_at_3",
            "strategy_mrr",
            "mode",
            "dialogue",
            "reason",
        ]
    )
    for row in detail_rows:
        ws_details.append(
            [
                row["couple"],
                row["sample_id"],
                row["turn_id"],
                row["row_index"],
                row["timestamp"],
                row["gold_target"],
                row["gold_target_source"],
                row["gold_strategy"],
                row["pred_strategy_top1"],
                row["pred_strategy_top2"],
                row["pred_strategy_top3"],
                row["strategy_gold_rank"],
                row["strategy_recall_at_1"],
                row["strategy_recall_at_2"],
                row["strategy_recall_at_3"],
                row["strategy_mrr"],
                row["mode"],
                row["dialogue"],
                row["reason"],
            ]
        )

    ws_conf = workbook.create_sheet("strategy_confusions")
    ws_conf.append(["gold_strategy", "pred_strategy_top1", "count"])
    for row in strategy_confusion_rows:
        ws_conf.append([row["gold_strategy"], row["pred_strategy_top1"], row["count"]])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_result_payload(
    *,
    mode: str,
    model: str,
    api_base: str,
    data_root: Path,
    context_size: int,
    include_target_internal_emotion: bool,
    detail_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    strategy_confusion_rows: Sequence[Dict[str, Any]],
) -> Dict[str, Any]:
    group_mode = clean_text(detail_rows[0].get("group_mode")) if detail_rows else "couple"
    return {
        "task": TASK_NAME,
        "group_mode": group_mode or "couple",
        "mode": mode,
        "prediction_model": model,
        "prediction_api_base": api_base,
        "data_root": str(data_root),
        "context_size": context_size,
        "include_target_internal_emotion": include_target_internal_emotion,
        "known_target_policy": "Gold support target is provided to the model as known_support_target; only strategy ranking is evaluated.",
        "taxonomy": STRATEGY_DEFINITIONS_BY_MODE[group_mode or "couple"],
        "overall": list(overall_rows),
        "couples": list(couple_rows),
        "samples": list(sample_rows),
        "details": list(detail_rows),
        "strategy_confusions": list(strategy_confusion_rows),
    }


def load_existing_details(
    output_json: Path,
    mode: str,
    model: str,
    context_size: int,
    group_mode: str,
    include_target_internal_emotion: bool,
) -> List[Dict[str, Any]]:
    if not output_json.exists():
        return []
    payload = json.loads(output_json.read_text(encoding="utf-8"))
    if clean_text(payload.get("task")) != TASK_NAME:
        return []
    if clean_text(payload.get("mode")) != mode:
        raise ValueError(f"Existing output_json mode={payload.get('mode')} != current mode={mode}")
    if clean_text(payload.get("prediction_model")) != model:
        raise ValueError(
            f"Existing output_json prediction_model={payload.get('prediction_model')} != current model={model}"
        )
    if int(payload.get("context_size") or 0) != context_size:
        raise ValueError(
            f"Existing output_json context_size={payload.get('context_size')} != current context_size={context_size}"
        )
    if bool(payload.get("include_target_internal_emotion")) != bool(include_target_internal_emotion):
        raise ValueError(
            "Existing output_json include_target_internal_emotion="
            f"{payload.get('include_target_internal_emotion')} != current include_target_internal_emotion={include_target_internal_emotion}"
        )
    if clean_text(payload.get("group_mode") or "couple") != group_mode:
        raise ValueError(
            f"Existing output_json group_mode={payload.get('group_mode')} != current group_mode={group_mode}"
        )
    return list(payload.get("details", []) or [])


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Evaluate next-turn therapist support-strategy ranking from prior context with known support target."
    )
    parser.add_argument("--data-root", default="data")
    parser.add_argument("--output-json", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--sample-ids-file", default="")
    parser.add_argument("--limit-samples", type=int, default=0)
    parser.add_argument("--context-size", type=int, default=3, help="Number of prior rows to include. Use 0 for none, -1 for all.")
    parser.add_argument(
        "--mode",
        choices=sorted(VALID_MODES | set(MODE_ALIASES)),
        default="dialogue",
    )
    parser.add_argument("--batch-size", type=int, default=3)
    parser.add_argument("--upload-cache", default=str(DEFAULT_UPLOAD_CACHE))
    parser.add_argument(
        "--api-base",
        default=os.getenv("BAILIAN_API_BASE")
        or os.getenv("DASHSCOPE_API_BASE")
        or os.getenv("OPENAI_API_BASE")
        or os.getenv("OPENAI_BASE_URL")
        or DEFAULT_API_BASE,
    )
    parser.add_argument(
        "--api-key",
        default=os.getenv("BAILIAN_API_KEY") or os.getenv("DASHSCOPE_API_KEY") or os.getenv("OPENAI_API_KEY"),
    )
    parser.add_argument("--model", default=os.getenv("EVAL_MODEL") or os.getenv("BAILIAN_MODEL") or "deepseek-v4-flash")
    parser.add_argument("--max-tokens", type=int, default=2200)
    parser.add_argument("--disable-thinking", action="store_true")
    parser.add_argument(
        "--include-target-internal-emotion",
        action="store_true",
        help="Provide the most recent available internal_emotion annotation for the known support target within the context window.",
    )
    args = parser.parse_args()

    mode = MODE_ALIASES.get(args.mode, args.mode)
    if mode not in VALID_MODES:
        raise ValueError(f"Unsupported mode: {args.mode}")
    if args.context_size < -1:
        raise ValueError("context_size must be >= -1. Use -1 for all prior rows.")
    if args.batch_size <= 0:
        raise ValueError("batch_size must be >= 1.")
    if not args.api_key:
        raise ValueError("Missing API key.")
    if not args.model:
        raise ValueError("Missing model.")

    data_root = Path(args.data_root).expanduser().resolve()
    output_json = Path(args.output_json).expanduser().resolve()
    output_xlsx = Path(args.output_xlsx).expanduser().resolve()
    upload_cache_path = Path(args.upload_cache).expanduser().resolve()
    upload_cache: Dict[str, Any] = {}
    if upload_cache_path.exists():
        try:
            upload_cache = json.loads(upload_cache_path.read_text(encoding="utf-8"))
        except Exception:
            upload_cache = {}

    api_base = normalize_api_base(args.api_base)
    allowed_sample_ids = load_sample_id_filter(args.sample_ids_file)

    samples = load_data_samples(data_root, ROOT_DIR)
    if allowed_sample_ids:
        samples = [sample for sample in samples if sample["sample_id"] in allowed_sample_ids]
    if args.limit_samples > 0:
        samples = samples[: args.limit_samples]

    detail_rows = load_existing_details(
        output_json,
        mode,
        args.model,
        args.context_size,
        clean_text(samples[0].get("group_mode") or "couple") if samples else "couple",
        args.include_target_internal_emotion,
    )
    completed_sample_ids = {row["sample_id"] for row in build_sample_rows_from_details(detail_rows)}

    targets_by_sample: Dict[str, List[Dict[str, Any]]] = {}
    total_checkpoint_count = len(detail_rows)
    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            continue
        targets = build_targets(
            sample,
            context_size=args.context_size,
            require_video=(mode == "turn_video_audio"),
        )
        if targets:
            targets_by_sample[sample_id] = targets
            total_checkpoint_count += len(targets)

    total_progress = tqdm(
        total=total_checkpoint_count,
        initial=len(detail_rows),
        desc="strategy total",
        unit="turn",
        dynamic_ncols=True,
    )

    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            print(f"[skip] {sample_id}", flush=True)
            continue

        targets = targets_by_sample.get(sample_id, [])
        if not targets:
            continue

        pred_by_turn: Dict[str, Dict[str, Any]] = {}
        with tqdm(total=len(targets), desc=f"strategy {sample_id}", unit="turn", dynamic_ncols=True, disable=True) as progress:
            if mode == "turn_video_audio":
                for target in targets:
                    before = len(pred_by_turn)
                    pred = request_turn_video_prediction(
                        api_base=api_base,
                        api_key=args.api_key,
                        model=args.model,
                        target=target,
                        include_target_internal_emotion=args.include_target_internal_emotion,
                        upload_cache=upload_cache,
                        upload_cache_path=upload_cache_path,
                        context_size=args.context_size,
                        max_tokens=args.max_tokens,
                        disable_thinking=args.disable_thinking,
                    )
                    pred_by_turn[target["turn_id"]] = pred
                    delta = max(0, len(pred_by_turn) - before)
                    progress.update(delta)
                    total_progress.update(delta)
                    progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                    total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                    time.sleep(0.8)
            else:
                for batch in batched(targets, args.batch_size):
                    before = len(pred_by_turn)
                    records = [
                        strategy_prompt_record(
                            target,
                            mode,
                            args.context_size,
                            args.include_target_internal_emotion,
                        )
                        for target in batch
                    ]
                    payload = request_text_predictions(
                        api_base=api_base,
                        api_key=args.api_key,
                        model=args.model,
                        sample_id=sample_id,
                        group_mode=clean_text(sample.get("group_mode") or "couple"),
                        targets=records,
                        mode=mode,
                        context_size=args.context_size,
                        include_target_internal_emotion=args.include_target_internal_emotion,
                        max_tokens=args.max_tokens,
                        disable_thinking=args.disable_thinking,
                    )
                    raw_items = extract_prediction_items(payload)
                    for item in raw_items:
                        turn_id = clean_text(item.get("turn_id"))
                        if not turn_id and len(batch) == 1:
                            turn_id = batch[0]["turn_id"]
                        if turn_id:
                            pred_by_turn[turn_id] = item
                    delta = max(0, len(pred_by_turn) - before)
                    progress.update(delta)
                    total_progress.update(delta)
                    progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                    total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                    time.sleep(0.3)

            missing = [target["turn_id"] for target in targets if target["turn_id"] not in pred_by_turn]
            if missing and mode != "turn_video_audio":
                missing_targets = [target for target in targets if target["turn_id"] in set(missing)]
                for target in missing_targets:
                    before = len(pred_by_turn)
                    payload = request_text_predictions(
                        api_base=api_base,
                        api_key=args.api_key,
                        model=args.model,
                        sample_id=sample_id,
                        group_mode=clean_text(sample.get("group_mode") or "couple"),
                        targets=[
                            strategy_prompt_record(
                                target,
                                mode,
                                args.context_size,
                                args.include_target_internal_emotion,
                            )
                        ],
                        mode=mode,
                        context_size=args.context_size,
                        include_target_internal_emotion=args.include_target_internal_emotion,
                        max_tokens=min(900, args.max_tokens),
                        disable_thinking=args.disable_thinking,
                    )
                    raw_items = coerce_single_target_prediction_items(
                        payload,
                        target["turn_id"],
                        strategy_definitions_for_sample(sample),
                    )
                    for item in raw_items:
                        turn_id = clean_text(item.get("turn_id")) or target["turn_id"]
                        if turn_id == target["turn_id"]:
                            pred_by_turn[turn_id] = item
                    delta = max(0, len(pred_by_turn) - before)
                    progress.update(delta)
                    total_progress.update(delta)
                    progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                    total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                    time.sleep(0.3)
                missing = [target["turn_id"] for target in targets if target["turn_id"] not in pred_by_turn]
        if missing:
            raise RuntimeError(f"missing predictions for {missing}")

        sample_detail_rows: List[Dict[str, Any]] = []
        strategy_definitions = strategy_definitions_for_sample(sample)
        for target in targets:
            pred = pred_by_turn[target["turn_id"]]
            strategy_top3 = normalize_strategy_top3(pred.get("strategy_top3", []), strategy_definitions)
            metrics = rank_metrics(target["gold_strategy"], strategy_top3)
            row = {
                "couple": target["couple"],
                "sample_id": sample_id,
                "group_mode": clean_text(sample.get("group_mode") or "couple"),
                "turn_id": target["turn_id"],
                "row_index": target["row_index"],
                "timestamp": target["timestamp"],
                "gold_target": target["gold_target_label"],
                "gold_target_source": target["gold_target_source"],
                "gold_strategy": target["gold_strategy"],
                "pred_strategy_top1": strategy_top3[0],
                "pred_strategy_top2": strategy_top3[1],
                "pred_strategy_top3": strategy_top3[2],
                "strategy_gold_rank": metrics["gold_rank"],
                "strategy_recall_at_1": metrics["recall_at_1"],
                "strategy_recall_at_2": metrics["recall_at_2"],
                "strategy_recall_at_3": metrics["recall_at_3"],
                "strategy_mrr": metrics["mrr"],
                "mode": mode,
                "dialogue": target["dialogue"],
                "reason": clean_text(pred.get("reason")),
            }
            detail_rows.append(row)
            sample_detail_rows.append(row)

        sample_metrics = summarize_rows(sample_detail_rows)
        completed_sample_ids.add(sample_id)

        sample_rows = build_sample_rows_from_details(detail_rows)
        overall_rows = [{"scope": "overall", **summarize_rows(detail_rows)}]
        couple_rows = build_couple_rows_from_details(detail_rows)
        strategy_confusion_rows = build_strategy_confusion_rows(detail_rows)

        output_json.parent.mkdir(parents=True, exist_ok=True)
        output_json.write_text(
            json.dumps(
                build_result_payload(
                    mode=mode,
                    model=args.model,
                    api_base=api_base,
                    data_root=data_root,
                    context_size=args.context_size,
                    include_target_internal_emotion=args.include_target_internal_emotion,
                    detail_rows=detail_rows,
                    sample_rows=sample_rows,
                    overall_rows=overall_rows,
                    couple_rows=couple_rows,
                    strategy_confusion_rows=strategy_confusion_rows,
                ),
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        print(
            f"[done] {sample_id}: turns={sample_metrics['turn_count']} "
            f"r1={sample_metrics['strategy_recall_at_1']:.3f} "
            f"r2={sample_metrics['strategy_recall_at_2']:.3f} "
            f"r3={sample_metrics['strategy_recall_at_3']:.3f}",
            flush=True,
        )

    total_progress.close()

    overall_rows = [{"scope": "overall", **summarize_rows(detail_rows)}]
    sample_rows = build_sample_rows_from_details(detail_rows)
    couple_rows = build_couple_rows_from_details(detail_rows)
    strategy_confusion_rows = build_strategy_confusion_rows(detail_rows)

    result = build_result_payload(
        mode=mode,
        model=args.model,
        api_base=api_base,
        data_root=data_root,
        context_size=args.context_size,
        include_target_internal_emotion=args.include_target_internal_emotion,
        detail_rows=detail_rows,
        sample_rows=sample_rows,
        overall_rows=overall_rows,
        couple_rows=couple_rows,
        strategy_confusion_rows=strategy_confusion_rows,
    )
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    write_result_workbook(
        output_xlsx,
        overall_rows,
        couple_rows,
        sample_rows,
        detail_rows,
        strategy_confusion_rows,
    )
    print(f"[written] {output_json}", flush=True)
    print(f"[written] {output_xlsx}", flush=True)


if __name__ == "__main__":
    main()
