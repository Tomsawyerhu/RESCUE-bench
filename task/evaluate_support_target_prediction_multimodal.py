import argparse
import json
import os
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
import requests
try:
    from tqdm.auto import tqdm
except Exception:  # pragma: no cover - tqdm is optional for script progress.
    def tqdm(iterable=None, **kwargs):
        return iterable if iterable is not None else _NullProgress()


    class _NullProgress:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def update(self, n=1):
            return None

        def set_postfix_str(self, *args, **kwargs):
            return None


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from task.eval_common import (
    DEFAULT_API_BASE,
    call_json_model,
    canonical_target,
    clean_text,
    group_label_for_mode,
    load_data_samples,
    load_sample_id_filter,
    mask_sensitive_terms,
    normalize_api_base,
    ensure_unique_ordered,
    rank_metrics,
    target_label,
)
from task.multimodal_eval_common import (
    MODE_ALIASES,
    VALID_TEXT_MODES,
    build_prior_context_rows,
    load_review_manifest,
    prompt_record_for_mode,
    row_timestamp,
)
from utils.project_env import load_project_env


load_project_env(ROOT_DIR)


TASK_NAME = "predict_next_therapist_support_target_from_prior_context_only"


def batched(seq: Sequence[Dict[str, Any]], size: int) -> List[List[Dict[str, Any]]]:
    return [list(seq[i : i + size]) for i in range(0, len(seq), size)]


def normalize_target_top3(values: Sequence[Any], candidate_target_labels: Sequence[str]) -> List[str]:
    labels = [label for label in ensure_unique_ordered(values) if label in candidate_target_labels]
    while len(labels) < 3:
        labels.append("")
    return labels[:3]


def current_row_support_target(row: Dict[str, Any], group_mode: str) -> tuple[str, str]:
    items = row.get("support_strategy") or []
    if not items:
        return "", "none"
    first = items[0] if isinstance(items[0], dict) else {}
    support_type = clean_text(first.get("strategy_type"))
    if not support_type:
        return "", "none"
    target = canonical_target(first.get("target") or [])
    if target:
        return target_label(target, group_label=group_label_for_mode(group_mode)), "support_strategy"
    fallback_target = canonical_target(row.get("target") or [])
    if fallback_target:
        return target_label(fallback_target, group_label=group_label_for_mode(group_mode)), "row_target_fallback"
    return "", "none"


def build_targets(sample: Dict[str, Any], context_size: int) -> List[Dict[str, Any]]:
    rows = sample["rows"]
    manifest_rows = load_review_manifest(sample["sample_id"])
    targets: List[Dict[str, Any]] = []

    for row_index, row in enumerate(rows, start=1):
        if clean_text(row.get("primary_speaker")) != "Therapist":
            continue

        gold_target_label, gold_target_source = current_row_support_target(row, clean_text(sample.get("group_mode") or "couple"))
        if not gold_target_label:
            continue

        candidate_target_labels = [label for label in sample["candidate_target_labels"] if label]
        if gold_target_label not in candidate_target_labels:
            candidate_target_labels.append(gold_target_label)

        targets.append(
            {
                "sample_id": sample["sample_id"],
                "couple": sample["couple"],
                "file_rel": sample["file_rel"],
                "turn_id": "%s__row%d" % (Path(sample["file_rel"]).stem, row_index),
                "row_index": row_index,
                "timestamp": row_timestamp(row),
                "dialogue": clean_text(row.get("dialogue_cleaned")),
                "gold_target_label": gold_target_label,
                "gold_target_source": gold_target_source,
                "participants": list(sample["participants"]),
                "candidate_target_labels": candidate_target_labels,
                "prior_context": build_prior_context_rows(rows, manifest_rows, row_index, context_size),
            }
        )

    return targets


def target_prompt_record(target: Dict[str, Any], mode: str, context_size: int) -> Dict[str, Any]:
    record: Dict[str, Any] = {
        "turn_id": target["turn_id"],
        "row_index": target["row_index"],
        "participants": target["participants"],
        "candidate_target_labels": target["candidate_target_labels"],
    }
    if context_size != 0:
        record["prior_context"] = [prompt_record_for_mode(item, mode) for item in target["prior_context"]]
    else:
        record["prior_context"] = []
    return record


def sanitize_target_records(records: Sequence[Dict[str, Any]], level: str) -> List[Dict[str, Any]]:
    sanitized: List[Dict[str, Any]] = []
    for record in records:
        item = json.loads(json.dumps(record, ensure_ascii=False))
        if level == "none":
            sanitized.append(item)
            continue
        for ctx in item.get("prior_context", []) or []:
            if "dialogue_cleaned" in ctx:
                if level == "mask":
                    ctx["dialogue_cleaned"] = mask_sensitive_terms(clean_text(ctx.get("dialogue_cleaned")))
                else:
                    ctx["dialogue_cleaned"] = "[redacted sensitive prior dialogue]"
            for bg in ctx.get("background_dialogue", []) or []:
                if "content" in bg:
                    if level == "mask":
                        bg["content"] = mask_sensitive_terms(clean_text(bg.get("content")))
                    else:
                        bg["content"] = "[redacted sensitive background dialogue]"
        sanitized.append(item)
    return sanitized


def prediction_prompt(sample_id: str, mode: str, records: Sequence[Dict[str, Any]], context_size: int) -> str:
    group_mode = "family" if any("Family" in record.get("candidate_target_labels", []) for record in records) else "couple"
    group_label = group_label_for_mode(group_mode)
    session_phrase = "family-therapy interaction" if group_mode == "family" else "couples-therapy interaction"
    mode_guidance = {
        "dialogue": "Use only language-side evidence from prior_context: primary_speaker, target, dialogue_cleaned, and background_dialogue.",
        "audio": "Use prior_context language-side evidence together with tone_of_voice. Do not use visual annotations.",
        "visual": "Use prior_context language plus visual annotation evidence from body_posture, facial_expressions, self_directed_behavior, and interaction_behavior.",
        "audiovisual": "Use prior_context language, visual annotations, and tone_of_voice.",
    }
    context_guidance = (
        "No prior dialogue context is provided.\n"
        if context_size == 0
        else (
            "Use all available previous rows in prior_context. Do not use any future information from the unseen therapist turn.\n"
            if context_size < 0
            else f"Use at most the previous {context_size} rows in prior_context. Do not use any future information from the unseen therapist turn.\n"
        )
    )
    return (
        f"You are predicting the therapist's NEXT supportive target in a {session_phrase}.\n"
        "You only see PRIOR CONTEXT before the target therapist turn. You must not assume access to the upcoming therapist utterance.\n"
        "Your task is to rank the 3 most likely support targets from candidate_target_labels.\n"
        f"{context_guidance}"
        f"Mode: {mode}\n"
        f"{mode_guidance[mode]}\n"
        "Target-label rules:\n"
        "- Predict who the therapist's next supportive move is mainly directed toward.\n"
        "- Candidate labels are closed-set; do not invent new labels.\n"
        f"- Use `{group_label}` when the likely move is jointly directed to the whole participant group, a shared pattern, or a shared task.\n"
        f"- Prefer a specific person instead of `{group_label}` when the likely move primarily invites, protects, challenges, reassures, or redirects one participant.\n"
        "- Use relational flow, emotional activation, withdrawal/pursuit patterns, floor-taking, and the therapist's ongoing focus in prior_context.\n"
        "- If a likely move would restore balance or protect one person, rank that person highly even if both partners are involved in the topic.\n"
        "- Return exactly 3 distinct target labels.\n"
        "- Echo the exact Sample value and exact input turn_id values from Targets. Never copy placeholder or example IDs.\n"
        "Output JSON only with this schema:\n"
        "{"
        "\"sample_id\":\"<echo the exact Sample value below>\","
        "\"predictions\":["
        "{"
        "\"turn_id\":\"<echo one exact turn_id from Targets>\","
        f"\"target_top3\":[\"ParticipantA\",\"{group_label}\",\"ParticipantB\"],"
        "\"reason\":\"short English reason based only on the prior context\""
        "}"
        "]"
        "}\n"
        "No markdown. JSON only.\n"
        f"Sample: {sample_id}\n"
        f"Targets: {json.dumps(list(records), ensure_ascii=False)}"
    )


def request_text_predictions(
    *,
    api_base: str,
    api_key: str,
    model: str,
    sample_id: str,
    targets: Sequence[Dict[str, Any]],
    mode: str,
    context_size: int,
    max_tokens: int,
    disable_thinking: bool,
) -> Dict[str, Any]:
    prompts = [
        prediction_prompt(sample_id, mode, sanitize_target_records(targets, "none"), context_size),
        prediction_prompt(sample_id, mode, sanitize_target_records(targets, "mask"), context_size),
        prediction_prompt(sample_id, mode, sanitize_target_records(targets, "redact"), context_size),
    ]
    last_error: Optional[Exception] = None
    for prompt_index, prompt in enumerate(prompts):
        for attempt in range(1, 4):
            try:
                return call_json_model(
                    api_base=api_base,
                    api_key=api_key,
                    model=model,
                    prompt=prompt,
                    max_tokens=max_tokens,
                    disable_thinking=disable_thinking,
                )
            except requests.HTTPError as exc:
                last_error = exc
                if "data_inspection_failed" in str(exc) and prompt_index < len(prompts) - 1:
                    break
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
            except (requests.RequestException, OSError) as exc:
                last_error = exc
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
    if last_error is not None:
        raise last_error
    raise RuntimeError("Unknown error while requesting text predictions")


def extract_prediction_items(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    items = payload.get("predictions", [])
    if isinstance(items, dict):
        items = [items]
    if not isinstance(items, list):
        items = []
    items = [item for item in items if isinstance(item, dict)]
    if not items and isinstance(payload, dict) and payload.get("target_top3"):
        items = [payload]
    return items


def coerce_single_target_prediction_items(payload: Dict[str, Any], expected_turn_id: str) -> List[Dict[str, Any]]:
    items = extract_prediction_items(payload)
    if len(items) != 1:
        return items
    item = dict(items[0])
    if not ensure_unique_ordered(item.get("target_top3", [])):
        return items
    item["turn_id"] = expected_turn_id
    return [item]


def summarize_rows(rows: Sequence[Dict[str, Any]]) -> Dict[str, float]:
    count = len(rows)
    if count == 0:
        return {
            "turn_count": 0,
            "target_recall_at_1": 0.0,
            "target_recall_at_2": 0.0,
            "target_recall_at_3": 0.0,
            "target_mrr": 0.0,
        }
    return {
        "turn_count": count,
        "target_recall_at_1": sum(row["target_recall_at_1"] for row in rows) / count,
        "target_recall_at_2": sum(row["target_recall_at_2"] for row in rows) / count,
        "target_recall_at_3": sum(row["target_recall_at_3"] for row in rows) / count,
        "target_mrr": sum(row["target_mrr"] for row in rows) / count,
    }


def build_sample_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_sample: Dict[tuple[str, str], List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_sample[(row["couple"], row["sample_id"])].append(row)
    return [
        {"couple": couple, "sample_id": sample_id, **summarize_rows(rows)}
        for (couple, sample_id), rows in sorted(by_sample.items(), key=lambda item: item[0][1])
    ]


def build_couple_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_couple: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_couple[row["couple"]].append(row)
    return [{"scope": couple, **summarize_rows(rows)} for couple, rows in sorted(by_couple.items(), key=lambda item: item[0])]


def build_target_confusion_rows(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    counter: Counter = Counter()
    for row in detail_rows:
        counter[(row["gold_target"], row["pred_target_top1"])] += 1
    return [
        {"gold_target": gold, "pred_target_top1": pred, "count": count}
        for (gold, pred), count in counter.most_common()
    ]


def write_result_workbook(
    output_path: Path,
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    detail_rows: Sequence[Dict[str, Any]],
    target_confusion_rows: Sequence[Dict[str, Any]],
) -> None:
    workbook = Workbook()

    ws_overall = workbook.active
    ws_overall.title = "overall"
    ws_overall.append(
        ["scope", "turn_count", "target_recall_at_1", "target_recall_at_2", "target_recall_at_3", "target_mrr"]
    )
    for row in overall_rows:
        ws_overall.append(
            [
                row["scope"],
                row["turn_count"],
                row["target_recall_at_1"],
                row["target_recall_at_2"],
                row["target_recall_at_3"],
                row["target_mrr"],
            ]
        )

    ws_couples = workbook.create_sheet("couples")
    ws_couples.append(list(ws_overall.iter_rows(min_row=1, max_row=1, values_only=True))[0])
    for row in couple_rows:
        ws_couples.append(
            [
                row["scope"],
                row["turn_count"],
                row["target_recall_at_1"],
                row["target_recall_at_2"],
                row["target_recall_at_3"],
                row["target_mrr"],
            ]
        )

    ws_samples = workbook.create_sheet("samples")
    ws_samples.append(
        ["couple", "sample_id", "turn_count", "target_recall_at_1", "target_recall_at_2", "target_recall_at_3", "target_mrr"]
    )
    for row in sample_rows:
        ws_samples.append(
            [
                row["couple"],
                row["sample_id"],
                row["turn_count"],
                row["target_recall_at_1"],
                row["target_recall_at_2"],
                row["target_recall_at_3"],
                row["target_mrr"],
            ]
        )

    ws_details = workbook.create_sheet("details")
    ws_details.append(
        [
            "couple",
            "sample_id",
            "turn_id",
            "row_index",
            "timestamp",
            "gold_target",
            "gold_target_source",
            "pred_target_top1",
            "pred_target_top2",
            "pred_target_top3",
            "target_gold_rank",
            "target_recall_at_1",
            "target_recall_at_2",
            "target_recall_at_3",
            "target_mrr",
            "mode",
            "dialogue",
            "reason",
        ]
    )
    for row in detail_rows:
        ws_details.append(
            [
                row["couple"],
                row["sample_id"],
                row["turn_id"],
                row["row_index"],
                row["timestamp"],
                row["gold_target"],
                row["gold_target_source"],
                row["pred_target_top1"],
                row["pred_target_top2"],
                row["pred_target_top3"],
                row["target_gold_rank"],
                row["target_recall_at_1"],
                row["target_recall_at_2"],
                row["target_recall_at_3"],
                row["target_mrr"],
                row["mode"],
                row["dialogue"],
                row["reason"],
            ]
        )

    ws_conf = workbook.create_sheet("target_confusions")
    ws_conf.append(["gold_target", "pred_target_top1", "count"])
    for row in target_confusion_rows:
        ws_conf.append([row["gold_target"], row["pred_target_top1"], row["count"]])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_result_payload(
    *,
    mode: str,
    model: str,
    api_base: str,
    data_root: Path,
    context_size: int,
    detail_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    target_confusion_rows: Sequence[Dict[str, Any]],
) -> Dict[str, Any]:
    return {
        "task": TASK_NAME,
        "mode": mode,
        "prediction_model": model,
        "prediction_api_base": api_base,
        "data_root": str(data_root),
        "context_size": context_size,
        "overall": list(overall_rows),
        "couples": list(couple_rows),
        "samples": list(sample_rows),
        "details": list(detail_rows),
        "target_confusions": list(target_confusion_rows),
        "selection_policy": "Only therapist turns with non-empty support_strategy annotations are evaluated. The model predicts support target only.",
    }


def load_existing_details(output_json: Path, mode: str, model: str, context_size: int) -> List[Dict[str, Any]]:
    if not output_json.exists():
        return []
    payload = json.loads(output_json.read_text(encoding="utf-8"))
    if clean_text(payload.get("task")) != TASK_NAME:
        return []
    if clean_text(payload.get("mode")) != mode:
        raise ValueError(f"Existing output_json mode={payload.get('mode')} != current mode={mode}")
    if clean_text(payload.get("prediction_model")) != model:
        raise ValueError(
            f"Existing output_json prediction_model={payload.get('prediction_model')} != current model={model}"
        )
    if int(payload.get("context_size") or 0) != context_size:
        raise ValueError(
            f"Existing output_json context_size={payload.get('context_size')} != current context_size={context_size}"
        )
    return list(payload.get("details", []) or [])


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Evaluate next-turn therapist support-target ranking from prior context with four multimodal modes."
    )
    parser.add_argument("--data-root", default="data")
    parser.add_argument("--output-json", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--sample-ids-file", default="")
    parser.add_argument("--limit-samples", type=int, default=0)
    parser.add_argument("--context-size", type=int, default=3, help="Number of prior rows to include. Use 0 for none, -1 for all.")
    parser.add_argument("--mode", choices=sorted(VALID_TEXT_MODES | {"text"}), default="dialogue")
    parser.add_argument("--batch-size", type=int, default=3)
    parser.add_argument(
        "--api-base",
        default=os.getenv("BAILIAN_API_BASE")
        or os.getenv("DASHSCOPE_API_BASE")
        or os.getenv("OPENAI_API_BASE")
        or os.getenv("OPENAI_BASE_URL")
        or DEFAULT_API_BASE,
    )
    parser.add_argument(
        "--api-key",
        default=os.getenv("BAILIAN_API_KEY") or os.getenv("DASHSCOPE_API_KEY") or os.getenv("OPENAI_API_KEY"),
    )
    parser.add_argument("--model", default=os.getenv("EVAL_MODEL") or os.getenv("BAILIAN_MODEL") or "deepseek-v4-flash")
    parser.add_argument("--max-tokens", type=int, default=1800)
    parser.add_argument("--disable-thinking", action="store_true")
    args = parser.parse_args()

    mode = MODE_ALIASES.get(args.mode, args.mode)
    if mode not in VALID_TEXT_MODES:
        raise ValueError(f"Unsupported mode: {args.mode}")
    if args.context_size < -1:
        raise ValueError("context_size must be >= -1. Use -1 for all prior rows.")
    if args.batch_size <= 0:
        raise ValueError("batch_size must be >= 1.")
    if not args.api_key:
        raise ValueError("Missing API key.")
    if not args.model:
        raise ValueError("Missing model.")

    data_root = Path(args.data_root).expanduser().resolve()
    output_json = Path(args.output_json).expanduser().resolve()
    output_xlsx = Path(args.output_xlsx).expanduser().resolve()
    api_base = normalize_api_base(args.api_base)
    allowed_sample_ids = load_sample_id_filter(args.sample_ids_file)

    samples = load_data_samples(data_root, ROOT_DIR)
    if allowed_sample_ids:
        samples = [sample for sample in samples if sample["sample_id"] in allowed_sample_ids]
    if args.limit_samples > 0:
        samples = samples[: args.limit_samples]

    detail_rows = load_existing_details(output_json, mode, args.model, args.context_size)
    completed_sample_ids = {row["sample_id"] for row in build_sample_rows_from_details(detail_rows)}

    targets_by_sample: Dict[str, List[Dict[str, Any]]] = {}
    total_checkpoint_count = len(detail_rows)
    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            continue
        targets = build_targets(sample, args.context_size)
        if targets:
            targets_by_sample[sample_id] = targets
            total_checkpoint_count += len(targets)

    total_progress = tqdm(
        total=total_checkpoint_count,
        initial=len(detail_rows),
        desc="support_target total",
        unit="turn",
        dynamic_ncols=True,
    )

    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            print(f"[skip] {sample_id}", flush=True)
            continue

        targets = targets_by_sample.get(sample_id, [])
        if not targets:
            continue

        pred_by_turn: Dict[str, Dict[str, Any]] = {}
        with tqdm(total=len(targets), desc=f"support_target {sample_id}", unit="turn", dynamic_ncols=True, disable=True) as progress:
            for batch in batched(targets, args.batch_size):
                before = len(pred_by_turn)
                records = [target_prompt_record(target, mode, args.context_size) for target in batch]
                payload = request_text_predictions(
                    api_base=api_base,
                    api_key=args.api_key,
                    model=args.model,
                    sample_id=sample_id,
                    targets=records,
                    mode=mode,
                    context_size=args.context_size,
                    max_tokens=args.max_tokens,
                    disable_thinking=args.disable_thinking,
                )
                raw_items = extract_prediction_items(payload)
                for item in raw_items:
                    turn_id = clean_text(item.get("turn_id"))
                    if not turn_id and len(batch) == 1:
                        turn_id = batch[0]["turn_id"]
                    if turn_id:
                        pred_by_turn[turn_id] = item
                delta = max(0, len(pred_by_turn) - before)
                progress.update(delta)
                total_progress.update(delta)
                progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                time.sleep(0.3)

            missing = [target["turn_id"] for target in targets if target["turn_id"] not in pred_by_turn]
            if missing:
                missing_targets = [target for target in targets if target["turn_id"] in set(missing)]
                for target in missing_targets:
                    before = len(pred_by_turn)
                    payload = request_text_predictions(
                        api_base=api_base,
                        api_key=args.api_key,
                        model=args.model,
                        sample_id=sample_id,
                        targets=[target_prompt_record(target, mode, args.context_size)],
                        mode=mode,
                        context_size=args.context_size,
                        max_tokens=min(900, args.max_tokens),
                        disable_thinking=args.disable_thinking,
                    )
                    raw_items = coerce_single_target_prediction_items(payload, target["turn_id"])
                    for item in raw_items:
                        turn_id = clean_text(item.get("turn_id")) or target["turn_id"]
                        if turn_id == target["turn_id"]:
                            pred_by_turn[turn_id] = item
                    delta = max(0, len(pred_by_turn) - before)
                    progress.update(delta)
                    total_progress.update(delta)
                    progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                    total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                    time.sleep(0.3)
                missing = [target["turn_id"] for target in targets if target["turn_id"] not in pred_by_turn]
        if missing:
            raise RuntimeError(f"missing predictions for {missing}")

        sample_detail_rows: List[Dict[str, Any]] = []
        for target in targets:
            pred = pred_by_turn[target["turn_id"]]
            target_top3 = normalize_target_top3(pred.get("target_top3", []), target["candidate_target_labels"])
            metrics = rank_metrics(target["gold_target_label"], target_top3)
            row = {
                "couple": target["couple"],
                "sample_id": sample_id,
                "turn_id": target["turn_id"],
                "row_index": target["row_index"],
                "timestamp": target["timestamp"],
                "gold_target": target["gold_target_label"],
                "gold_target_source": target["gold_target_source"],
                "pred_target_top1": target_top3[0],
                "pred_target_top2": target_top3[1],
                "pred_target_top3": target_top3[2],
                "target_gold_rank": metrics["gold_rank"],
                "target_recall_at_1": metrics["recall_at_1"],
                "target_recall_at_2": metrics["recall_at_2"],
                "target_recall_at_3": metrics["recall_at_3"],
                "target_mrr": metrics["mrr"],
                "mode": mode,
                "dialogue": target["dialogue"],
                "reason": clean_text(pred.get("reason")),
            }
            detail_rows.append(row)
            sample_detail_rows.append(row)

        sample_metrics = summarize_rows(sample_detail_rows)
        completed_sample_ids.add(sample_id)

        sample_rows = build_sample_rows_from_details(detail_rows)
        overall_rows = [{"scope": "overall", **summarize_rows(detail_rows)}]
        couple_rows = build_couple_rows_from_details(detail_rows)
        target_confusion_rows = build_target_confusion_rows(detail_rows)

        output_json.parent.mkdir(parents=True, exist_ok=True)
        output_json.write_text(
            json.dumps(
                build_result_payload(
                    mode=mode,
                    model=args.model,
                    api_base=api_base,
                    data_root=data_root,
                    context_size=args.context_size,
                    detail_rows=detail_rows,
                    sample_rows=sample_rows,
                    overall_rows=overall_rows,
                    couple_rows=couple_rows,
                    target_confusion_rows=target_confusion_rows,
                ),
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        print(
            f"[done] {sample_id}: turns={sample_metrics['turn_count']} "
            f"r1={sample_metrics['target_recall_at_1']:.3f} "
            f"r2={sample_metrics['target_recall_at_2']:.3f} "
            f"r3={sample_metrics['target_recall_at_3']:.3f}",
            flush=True,
        )

    total_progress.close()

    overall_rows = [{"scope": "overall", **summarize_rows(detail_rows)}]
    sample_rows = build_sample_rows_from_details(detail_rows)
    couple_rows = build_couple_rows_from_details(detail_rows)
    target_confusion_rows = build_target_confusion_rows(detail_rows)

    result = build_result_payload(
        mode=mode,
        model=args.model,
        api_base=api_base,
        data_root=data_root,
        context_size=args.context_size,
        detail_rows=detail_rows,
        sample_rows=sample_rows,
        overall_rows=overall_rows,
        couple_rows=couple_rows,
        target_confusion_rows=target_confusion_rows,
    )
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    write_result_workbook(output_xlsx, overall_rows, couple_rows, sample_rows, detail_rows, target_confusion_rows)
    print(f"[written] {output_json}", flush=True)
    print(f"[written] {output_xlsx}", flush=True)


if __name__ == "__main__":
    main()
