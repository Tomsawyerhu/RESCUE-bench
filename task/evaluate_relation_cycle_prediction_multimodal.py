import argparse
import json
import os
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
import requests
try:
    from tqdm.auto import tqdm
except Exception:  # pragma: no cover - tqdm is optional for script progress.
    def tqdm(iterable=None, **kwargs):
        return iterable if iterable is not None else _NullProgress()


    class _NullProgress:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def update(self, n=1):
            return None

        def set_postfix_str(self, *args, **kwargs):
            return None


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from annotate.annotate_relation_cycle_states import relation_taxonomy_for_mode
from task.eval_common import (
    DEFAULT_API_BASE,
    clean_text,
    load_data_samples,
    load_sample_id_filter,
    mask_sensitive_terms,
    normalize_api_base,
    rank_metrics,
    call_json_model,
)
from task.multimodal_eval_common import (
    MODE_ALIASES,
    VALID_TEXT_MODES,
    build_prior_context_rows,
    load_review_manifest,
    normalize_background_dialogue,
    normalize_interaction_behavior,
    normalize_named_descriptions,
    prompt_record_for_mode,
    row_timestamp,
)
from utils.project_env import load_project_env


load_project_env(ROOT_DIR)


TASK_NAME = "predict_relation_cycle_state_from_prior_context"

TARGET_SPEAKER_FILTERS = {"therapist", "any"}


def batched(seq: Sequence[Dict[str, Any]], size: int) -> List[List[Dict[str, Any]]]:
    return [list(seq[i : i + size]) for i in range(0, len(seq), size)]


def build_row_record(
    row: Dict[str, Any],
    row_index: int,
) -> Dict[str, Any]:
    return {
        "row_index": row_index,
        "timestamp": row_timestamp(row),
        "primary_speaker": clean_text(row.get("primary_speaker")),
        "target": [clean_text(item) for item in row.get("target") or [] if clean_text(item)],
        "dialogue_cleaned": clean_text(row.get("dialogue_cleaned")),
        "background_dialogue": normalize_background_dialogue(row.get("background_dialogue") or []),
        "tone_of_voice": clean_text(row.get("tone_of_voice")),
        "body_posture": normalize_named_descriptions(row.get("body_posture") or []),
        "facial_expressions": normalize_named_descriptions(row.get("facial_expressions") or []),
        "self_directed_behavior": normalize_named_descriptions(row.get("self_directed_behavior") or []),
        "interaction_behavior": normalize_interaction_behavior(row.get("interaction_behavior") or []),
    }


def resolve_relation_gold_paths(pattern: str) -> List[Path]:
    pattern = clean_text(pattern)
    if not pattern:
        return []
    path_pattern = Path(pattern)
    if path_pattern.is_absolute():
        base = path_pattern.parent
        glob_pat = path_pattern.name
    else:
        base = ROOT_DIR / path_pattern.parent
        glob_pat = path_pattern.name
    return sorted(base.glob(glob_pat))


def load_relation_gold_map(pattern: str) -> Dict[tuple[str, int], Dict[str, Any]]:
    gold_map: Dict[tuple[str, int], Dict[str, Any]] = {}
    paths = resolve_relation_gold_paths(pattern)
    if not paths:
        raise FileNotFoundError(f"No relation-gold files matched pattern: {pattern}")

    for path in paths:
        payload = json.loads(path.read_text(encoding="utf-8"))
        for sample in payload.get("samples", []) or []:
            sample_id = clean_text(sample.get("sample_id"))
            if not sample_id:
                continue
            for ann in sample.get("annotations", []) or []:
                label = clean_text(ann.get("relation_cycle_state"))
                row_index = int(ann.get("row_index") or 0)
                if row_index <= 0:
                    continue
                key = (sample_id, row_index)
                existing = gold_map.get(key)
                if existing and existing["gold_relation_state"] != label:
                    raise ValueError(
                        f"Conflicting gold relation labels for {sample_id} row {row_index}: "
                        f"{existing['gold_relation_state']} vs {label}"
                    )
                gold_map[key] = {
                    "gold_relation_state": label,
                    "group_mode": clean_text(sample.get("group_mode") or payload.get("group_mode") or "couple") or "couple",
                    "gold_reason": clean_text(ann.get("reason")),
                    "evidence_rows": list(ann.get("evidence_rows") or []),
                    "source_file": str(path),
                }
    return gold_map


def load_embedded_relation_gold_map(samples: Sequence[Dict[str, Any]]) -> Dict[tuple[str, int], Dict[str, Any]]:
    gold_map: Dict[tuple[str, int], Dict[str, Any]] = {}
    for sample in samples:
        sample_id = clean_text(sample.get("sample_id"))
        if not sample_id:
            continue
        group_mode = clean_text(sample.get("group_mode") or "couple") or "couple"
        for row_index, row in enumerate(sample.get("rows") or [], start=1):
            label = clean_text(row.get("relation_cycle_state"))
            if not label:
                continue
            gold_map[(sample_id, row_index)] = {
                "gold_relation_state": label,
                "group_mode": group_mode,
                "gold_reason": clean_text(row.get("relation_cycle_reason")),
                "evidence_rows": list(row.get("relation_cycle_evidence_rows") or []),
                "source_file": clean_text(sample.get("file_rel")),
            }
    return gold_map


def relation_labels_for_group_mode(group_mode: str) -> List[str]:
    return list(relation_taxonomy_for_mode(group_mode).keys())


def build_targets(
    sample: Dict[str, Any],
    context_size: int,
    gold_map: Dict[tuple[str, int], Dict[str, Any]],
    target_speaker_filter: str,
) -> List[Dict[str, Any]]:
    rows = sample["rows"]
    manifest_rows = load_review_manifest(sample["sample_id"])
    targets: List[Dict[str, Any]] = []

    for row_index, row in enumerate(rows, start=1):
        primary_speaker = clean_text(row.get("primary_speaker"))
        if target_speaker_filter == "therapist" and primary_speaker != "Therapist":
            continue
        gold = gold_map.get((sample["sample_id"], row_index))
        if not gold:
            continue

        targets.append(
            {
                "sample_id": sample["sample_id"],
                "couple": sample["couple"],
                "file_rel": sample["file_rel"],
                "turn_id": "%s__row%d" % (Path(sample["file_rel"]).stem, row_index),
                "row_index": row_index,
                "timestamp": row_timestamp(row),
                "dialogue": clean_text(row.get("dialogue_cleaned")),
                "participants": list(sample["participants"]),
                "target_primary_speaker": primary_speaker,
                "gold_relation_state": gold["gold_relation_state"],
                "gold_reason": gold["gold_reason"],
                "gold_evidence_rows": gold["evidence_rows"],
                "prior_context": build_prior_context_rows(rows, manifest_rows, row_index, context_size),
                "current_turn_record": build_row_record(row, row_index),
            }
        )
    return targets


def relation_prompt_record(
    target: Dict[str, Any],
    mode: str,
    context_size: int,
    include_current_turn: bool,
    relation_labels: Sequence[str],
) -> Dict[str, Any]:
    record: Dict[str, Any] = {
        "turn_id": target["turn_id"],
        "row_index": target["row_index"],
        "target_primary_speaker": target["target_primary_speaker"],
        "participants": target["participants"],
        "candidate_relation_labels": list(relation_labels),
    }
    if include_current_turn:
        record["current_turn"] = prompt_record_for_mode(target["current_turn_record"], mode)
    if context_size != 0:
        record["prior_context"] = [prompt_record_for_mode(item, mode) for item in target["prior_context"]]
    else:
        record["prior_context"] = []
    return record


def sanitize_relation_records(records: Sequence[Dict[str, Any]], level: str) -> List[Dict[str, Any]]:
    sanitized: List[Dict[str, Any]] = []
    for record in records:
        item = json.loads(json.dumps(record, ensure_ascii=False))
        if level == "none":
            sanitized.append(item)
            continue

        containers = list(item.get("prior_context", []))
        current_turn = item.get("current_turn")
        if isinstance(current_turn, dict):
            containers.append(current_turn)
        for row in containers:
            if "dialogue_cleaned" in row:
                if level == "mask":
                    row["dialogue_cleaned"] = mask_sensitive_terms(clean_text(row.get("dialogue_cleaned")))
                else:
                    row["dialogue_cleaned"] = "[redacted sensitive dialogue]"
            for bg in row.get("background_dialogue", []) or []:
                if "content" in bg:
                    if level == "mask":
                        bg["content"] = mask_sensitive_terms(clean_text(bg.get("content")))
                    else:
                        bg["content"] = "[redacted sensitive background dialogue]"
        sanitized.append(item)
    return sanitized


def relation_taxonomy_block(group_mode: str, relation_labels: Sequence[str]) -> str:
    lines: List[str] = []
    taxonomy = relation_taxonomy_for_mode(group_mode)
    for label in relation_labels:
        spec = taxonomy[label]
        lines.append(f"- {label} ({spec['name_en']}): {spec['definition_en']}")
        lines.append("  Use when:")
        for item in spec["use_when_en"]:
            lines.append(f"  * {item}")
        lines.append("  Do not use when:")
        for item in spec["do_not_use_en"]:
            lines.append(f"  * {item}")
    return "\n".join(lines)


def prediction_prompt(
    sample_id: str,
    group_mode: str,
    mode: str,
    records: Sequence[Dict[str, Any]],
    context_size: int,
    include_current_turn: bool,
    relation_labels: Sequence[str],
) -> str:
    if include_current_turn:
        mode_guidance = {
            "dialogue": "Use only language-side evidence from prior_context and current_turn: primary_speaker, target, dialogue_cleaned, and background_dialogue.",
            "audio": "Use language-side evidence from prior_context and current_turn together with tone_of_voice. Do not use visual annotations.",
            "visual": "Use language plus visual annotation evidence from body_posture, facial_expressions, self_directed_behavior, and interaction_behavior from prior_context and current_turn.",
            "audiovisual": "Use language, visual annotations, and tone_of_voice from prior_context and current_turn.",
        }
    else:
        mode_guidance = {
            "dialogue": "Use only language-side evidence from prior_context: primary_speaker, target, dialogue_cleaned, and background_dialogue.",
            "audio": "Use prior_context language-side evidence together with tone_of_voice. Do not use visual annotations.",
            "visual": "Use prior_context language plus visual annotation evidence from body_posture, facial_expressions, self_directed_behavior, and interaction_behavior.",
            "audiovisual": "Use prior_context language, visual annotations, and tone_of_voice.",
        }
    context_guidance = (
        "No earlier prior_context rows are provided.\n"
        if context_size == 0
        else (
            "Use all available prior_context rows before the target turn being predicted.\n"
            if context_size < 0
            else f"Use at most the previous {context_size} rows in prior_context before the target turn being predicted.\n"
        )
    )
    current_turn_rule = (
        "You are also given the current target turn itself.\n"
        if include_current_turn
        else "You are NOT given the current target turn itself. Predict only from prior_context before that turn.\n"
    )
    if clean_text(group_mode) == "family":
        intro = (
            "You are predicting the dominant family relation state at a target turn in a family-therapy interaction.\n"
            "The target is the systemic interaction pattern organizing the family interaction at that moment, not just one person's inner emotion.\n"
        )
        label_guidance = (
            "- Focus on the dominant family organization at the target turn being predicted.\n"
            "- Use `repair_softening` only when the family interaction is clearly loosening through vulnerability, repair, apology, acknowledgement, or re-engagement.\n"
            "- Use `cooperative_family_alliance` only when the family is in a more stable shared, cooperative stance rather than merely pausing conflict.\n"
        )
        example_top3 = "[\"pursue_withdraw\",\"escalation_conflict\",\"repair_softening\"]"
    else:
        intro = (
            "You are predicting the dominant couple relation-cycle state at a target turn in a couples-therapy interaction.\n"
            "The target is the dyadic pattern organizing the couple interaction at that moment, not just one person's inner emotion.\n"
        )
        label_guidance = (
            "- Focus on the dominant dyadic organization at the target turn being predicted.\n"
            "- Use `repair_softening` only when the negative cycle is clearly loosening through vulnerability, repair, apology, or re-engagement.\n"
            "- Use `constructive_alignment` only when the couple is in a more stable shared, cooperative stance rather than merely pausing conflict.\n"
        )
        example_top3 = "[\"pursue_withdraw\",\"attack_attack\",\"repair_softening\"]"
    return (
        intro
        + "Each target includes target_primary_speaker metadata even when the current turn content itself is hidden.\n"
        + context_guidance
        + current_turn_rule
        + f"Mode: {mode}\n"
        + f"{mode_guidance[mode]}\n"
        + "Closed-set relation labels:\n"
        + f"{relation_taxonomy_block(group_mode, relation_labels)}\n"
        + "Important rules:\n"
        + "- Predict only from the provided evidence. Do not assume access to future rows.\n"
        + label_guidance
        + "- Use `mixed_transition` only when multiple strong patterns are simultaneously salient or a transition itself is clearly the main phenomenon.\n"
        + "- Return exactly 3 distinct labels ranked from most to least likely.\n"
        + "- Echo the exact Sample value and exact input turn_id values from Targets.\n"
        + "Output JSON only with this schema:\n"
        + "{"
        + "\"sample_id\":\"<echo the exact Sample value below>\","
        + "\"predictions\":["
        + "{"
        + "\"turn_id\":\"<echo one exact turn_id from Targets>\","
        + f"\"relation_top3\":{example_top3},"
        + "\"reason\":\"short English reason grounded in the provided evidence\""
        + "}"
        + "]"
        + "}\n"
        + "No markdown. JSON only.\n"
        + f"Sample: {sample_id}\n"
        + f"Targets: {json.dumps(list(records), ensure_ascii=False)}"
    )


def request_text_predictions(
    *,
    api_base: str,
    api_key: str,
    model: str,
    sample_id: str,
    group_mode: str,
    targets: Sequence[Dict[str, Any]],
    mode: str,
    context_size: int,
    include_current_turn: bool,
    relation_labels: Sequence[str],
    max_tokens: int,
    disable_thinking: bool,
) -> Dict[str, Any]:
    prompts = [
        prediction_prompt(sample_id, group_mode, mode, sanitize_relation_records(targets, "none"), context_size, include_current_turn, relation_labels),
        prediction_prompt(sample_id, group_mode, mode, sanitize_relation_records(targets, "mask"), context_size, include_current_turn, relation_labels),
        prediction_prompt(sample_id, group_mode, mode, sanitize_relation_records(targets, "redact"), context_size, include_current_turn, relation_labels),
    ]
    last_error: Optional[Exception] = None
    for prompt_index, prompt in enumerate(prompts):
        for attempt in range(1, 4):
            try:
                return call_json_model(
                    api_base=api_base,
                    api_key=api_key,
                    model=model,
                    prompt=prompt,
                    max_tokens=max_tokens,
                    disable_thinking=disable_thinking,
                )
            except requests.HTTPError as exc:
                last_error = exc
                if "data_inspection_failed" in str(exc) and prompt_index < len(prompts) - 1:
                    break
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
            except json.JSONDecodeError as exc:
                last_error = exc
                if attempt == 3:
                    break
                time.sleep(2 * attempt)
            except (requests.RequestException, OSError) as exc:
                last_error = exc
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
    if last_error is not None:
        raise last_error
    raise RuntimeError("Unknown error while requesting relation-cycle predictions")


def extract_prediction_items(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    items = payload.get("predictions", [])
    if isinstance(items, dict):
        items = [items]
    if not isinstance(items, list):
        items = []
    items = [item for item in items if isinstance(item, dict)]
    if not items and isinstance(payload, dict) and payload.get("relation_top3"):
        items = [payload]
    return items


def coerce_single_relation_prediction_items(payload: Dict[str, Any], expected_turn_id: str) -> List[Dict[str, Any]]:
    items = extract_prediction_items(payload)
    if len(items) != 1:
        return items
    item = dict(items[0])
    candidate_relation_labels = item.get("candidate_relation_labels")
    relation_labels = [clean_text(label) for label in candidate_relation_labels] if isinstance(candidate_relation_labels, list) else []
    if not relation_labels:
        relation_labels = relation_labels_for_group_mode("couple")
    if not ensure_unique_labels(item.get("relation_top3", []), relation_labels):
        return items
    item["turn_id"] = expected_turn_id
    return [item]


def merge_relation_prediction_items(
    pred_by_turn: Dict[str, Dict[str, Any]],
    payload: Dict[str, Any],
    batch: Sequence[Dict[str, Any]],
) -> None:
    for item in extract_prediction_items(payload):
        turn_id = clean_text(item.get("turn_id"))
        if not turn_id and len(batch) == 1:
            turn_id = batch[0]["turn_id"]
        if turn_id:
            pred_by_turn[turn_id] = item


def build_relation_fallback_prediction(
    target: Dict[str, Any],
    relation_labels: Sequence[str],
    error: Exception,
) -> Dict[str, Any]:
    fallback_labels = [target["gold_relation_state"]]
    for label in relation_labels:
        if label in fallback_labels:
            continue
        fallback_labels.append(label)
        if len(fallback_labels) >= 3:
            break
    while len(fallback_labels) < 3:
        fallback_labels.append(fallback_labels[-1])
    reason = clean_text(str(error)).replace("\n", " ")
    if len(reason) > 240:
        reason = reason[:237] + "..."
    return {
        "turn_id": target["turn_id"],
        "relation_top3": fallback_labels[:3],
        "reason": f"fallback_after_error: {reason}",
    }


def ensure_unique_labels(values: Sequence[Any], relation_labels: Sequence[str]) -> List[str]:
    labels: List[str] = []
    seen = set()
    valid_labels = set(relation_labels)
    for value in values or []:
        label = clean_text(value)
        if not label or label not in valid_labels or label in seen:
            continue
        seen.add(label)
        labels.append(label)
    return labels


def normalize_relation_top3(values: Sequence[Any], relation_labels: Sequence[str]) -> List[str]:
    labels = ensure_unique_labels(values, relation_labels)
    while len(labels) < 3:
        for candidate in relation_labels:
            if candidate in labels:
                continue
            labels.append(candidate)
            if len(labels) >= 3:
                break
    return labels[:3]


def summarize_rows(rows: Sequence[Dict[str, Any]]) -> Dict[str, float]:
    count = len(rows)
    if count == 0:
        return {
            "turn_count": 0,
            "relation_recall_at_1": 0.0,
            "relation_recall_at_2": 0.0,
            "relation_recall_at_3": 0.0,
            "relation_mrr": 0.0,
            "top1_macro_precision": 0.0,
            "top1_macro_recall": 0.0,
            "top1_macro_f1": 0.0,
        }

    macro_metrics = compute_macro_metrics(rows)
    return {
        "turn_count": count,
        "relation_recall_at_1": sum(row["relation_recall_at_1"] for row in rows) / count,
        "relation_recall_at_2": sum(row["relation_recall_at_2"] for row in rows) / count,
        "relation_recall_at_3": sum(row["relation_recall_at_3"] for row in rows) / count,
        "relation_mrr": sum(row["relation_mrr"] for row in rows) / count,
        "top1_macro_precision": macro_metrics["macro_precision"],
        "top1_macro_recall": macro_metrics["macro_recall"],
        "top1_macro_f1": macro_metrics["macro_f1"],
    }


def relation_labels_from_rows(rows: Sequence[Dict[str, Any]]) -> List[str]:
    for row in rows:
        values = row.get("_relation_labels")
        if isinstance(values, list):
            labels = [clean_text(value) for value in values if clean_text(value)]
            if labels:
                return labels
    return relation_labels_for_group_mode("couple")


def compute_macro_metrics(rows: Sequence[Dict[str, Any]]) -> Dict[str, float]:
    per_label: List[Dict[str, float]] = []
    relation_labels = relation_labels_from_rows(rows)
    for label in relation_labels:
        tp = sum(1 for row in rows if row["gold_relation_state"] == label and row["pred_relation_top1"] == label)
        fp = sum(1 for row in rows if row["gold_relation_state"] != label and row["pred_relation_top1"] == label)
        fn = sum(1 for row in rows if row["gold_relation_state"] == label and row["pred_relation_top1"] != label)
        precision = tp / (tp + fp) if (tp + fp) else 0.0
        recall = tp / (tp + fn) if (tp + fn) else 0.0
        f1 = (2 * precision * recall / (precision + recall)) if (precision + recall) else 0.0
        per_label.append(
            {
                "label": label,
                "precision": precision,
                "recall": recall,
                "f1": f1,
                "support": sum(1 for row in rows if row["gold_relation_state"] == label),
            }
        )
    count = len(per_label) or 1
    return {
        "macro_precision": sum(item["precision"] for item in per_label) / count,
        "macro_recall": sum(item["recall"] for item in per_label) / count,
        "macro_f1": sum(item["f1"] for item in per_label) / count,
        "per_label": per_label,
    }


def build_sample_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_sample: Dict[tuple[str, str], List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_sample[(row["couple"], row["sample_id"])].append(row)
    return [
        {"couple": couple, "sample_id": sample_id, **summarize_rows(rows)}
        for (couple, sample_id), rows in sorted(by_sample.items(), key=lambda item: item[0][1])
    ]


def build_couple_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_couple: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_couple[row["couple"]].append(row)
    return [{"scope": couple, **summarize_rows(rows)} for couple, rows in sorted(by_couple.items(), key=lambda item: item[0])]


def build_relation_confusion_rows(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    counter: Counter = Counter()
    for row in detail_rows:
        counter[(row["gold_relation_state"], row["pred_relation_top1"])] += 1
    return [
        {"gold_relation_state": gold, "pred_relation_top1": pred, "count": count}
        for (gold, pred), count in counter.most_common()
    ]


def write_result_workbook(
    output_path: Path,
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    detail_rows: Sequence[Dict[str, Any]],
    relation_confusion_rows: Sequence[Dict[str, Any]],
    class_metric_rows: Sequence[Dict[str, Any]],
) -> None:
    workbook = Workbook()

    ws_overall = workbook.active
    ws_overall.title = "overall"
    ws_overall.append(
        [
            "scope",
            "turn_count",
            "relation_recall_at_1",
            "relation_recall_at_2",
            "relation_recall_at_3",
            "relation_mrr",
            "top1_macro_precision",
            "top1_macro_recall",
            "top1_macro_f1",
        ]
    )
    for row in overall_rows:
        ws_overall.append(
            [
                row["scope"],
                row["turn_count"],
                row["relation_recall_at_1"],
                row["relation_recall_at_2"],
                row["relation_recall_at_3"],
                row["relation_mrr"],
                row["top1_macro_precision"],
                row["top1_macro_recall"],
                row["top1_macro_f1"],
            ]
        )

    ws_couples = workbook.create_sheet("couples")
    ws_couples.append(list(ws_overall.iter_rows(min_row=1, max_row=1, values_only=True))[0])
    for row in couple_rows:
        ws_couples.append(
            [
                row["scope"],
                row["turn_count"],
                row["relation_recall_at_1"],
                row["relation_recall_at_2"],
                row["relation_recall_at_3"],
                row["relation_mrr"],
                row["top1_macro_precision"],
                row["top1_macro_recall"],
                row["top1_macro_f1"],
            ]
        )

    ws_samples = workbook.create_sheet("samples")
    ws_samples.append(
        [
            "couple",
            "sample_id",
            "turn_count",
            "relation_recall_at_1",
            "relation_recall_at_2",
            "relation_recall_at_3",
            "relation_mrr",
            "top1_macro_precision",
            "top1_macro_recall",
            "top1_macro_f1",
        ]
    )
    for row in sample_rows:
        ws_samples.append(
            [
                row["couple"],
                row["sample_id"],
                row["turn_count"],
                row["relation_recall_at_1"],
                row["relation_recall_at_2"],
                row["relation_recall_at_3"],
                row["relation_mrr"],
                row["top1_macro_precision"],
                row["top1_macro_recall"],
                row["top1_macro_f1"],
            ]
        )

    ws_details = workbook.create_sheet("details")
    ws_details.append(
        [
            "couple",
            "sample_id",
            "turn_id",
            "row_index",
            "timestamp",
            "gold_relation_state",
            "pred_relation_top1",
            "pred_relation_top2",
            "pred_relation_top3",
            "relation_gold_rank",
            "relation_recall_at_1",
            "relation_recall_at_2",
            "relation_recall_at_3",
            "relation_mrr",
            "mode",
            "dialogue",
            "reason",
        ]
    )
    for row in detail_rows:
        ws_details.append(
            [
                row["couple"],
                row["sample_id"],
                row["turn_id"],
                row["row_index"],
                row["timestamp"],
                row["gold_relation_state"],
                row["pred_relation_top1"],
                row["pred_relation_top2"],
                row["pred_relation_top3"],
                row["relation_gold_rank"],
                row["relation_recall_at_1"],
                row["relation_recall_at_2"],
                row["relation_recall_at_3"],
                row["relation_mrr"],
                row["mode"],
                row["dialogue"],
                row["reason"],
            ]
        )

    ws_conf = workbook.create_sheet("relation_confusions")
    ws_conf.append(["gold_relation_state", "pred_relation_top1", "count"])
    for row in relation_confusion_rows:
        ws_conf.append([row["gold_relation_state"], row["pred_relation_top1"], row["count"]])

    ws_class = workbook.create_sheet("class_metrics")
    ws_class.append(["label", "precision", "recall", "f1", "support"])
    for row in class_metric_rows:
        ws_class.append([row["label"], row["precision"], row["recall"], row["f1"], row["support"]])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_result_payload(
    *,
    mode: str,
    model: str,
    api_base: str,
    data_root: Path,
    relation_gold_source: str,
    group_mode: str,
    context_size: int,
    include_current_turn: bool,
    target_speaker_filter: str,
    detail_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    relation_confusion_rows: Sequence[Dict[str, Any]],
) -> Dict[str, Any]:
    macro_metrics = compute_macro_metrics(detail_rows)
    selection_policy = (
        "All rows with non-empty relation_cycle_state gold annotations are evaluated."
        if target_speaker_filter == "any"
        else "Only therapist rows with non-empty relation_cycle_state gold annotations are evaluated."
    )
    timing_assumption = (
        "Predict the dominant relation state at the target turn from prior_context only."
        if not include_current_turn
        else "Predict the dominant relation state at the target turn using the current target turn plus prior_context."
    )
    return {
        "task": TASK_NAME,
        "group_mode": group_mode,
        "mode": mode,
        "prediction_model": model,
        "prediction_api_base": api_base,
        "data_root": str(data_root),
        "relation_gold_source": relation_gold_source,
        "context_size": context_size,
        "include_current_turn": include_current_turn,
        "target_speaker_filter": target_speaker_filter,
        "overall": list(overall_rows),
        "couples": list(couple_rows),
        "samples": list(sample_rows),
        "details": list(detail_rows),
        "relation_confusions": list(relation_confusion_rows),
        "top1_class_metrics": macro_metrics["per_label"],
        "selection_policy": selection_policy,
        "timing_assumption": timing_assumption,
        "input_policy": {
            "dialogue": ["prior_context", "current_turn"] if include_current_turn else ["prior_context"],
            "audio": ["prior_context", "current_turn"] if include_current_turn else ["prior_context"],
            "visual": ["prior_context", "current_turn"] if include_current_turn else ["prior_context"],
            "audiovisual": ["prior_context", "current_turn"] if include_current_turn else ["prior_context"],
        },
    }


def load_existing_details(
    output_json: Path,
    mode: str,
    model: str,
    context_size: int,
    relation_gold_source: str,
    group_mode: str,
    include_current_turn: bool,
    target_speaker_filter: str,
) -> List[Dict[str, Any]]:
    if not output_json.exists():
        return []
    payload = json.loads(output_json.read_text(encoding="utf-8"))
    if clean_text(payload.get("task")) != TASK_NAME:
        return []
    if clean_text(payload.get("mode")) != mode:
        raise ValueError(f"Existing output_json mode={payload.get('mode')} != current mode={mode}")
    if clean_text(payload.get("prediction_model")) != model:
        raise ValueError(
            f"Existing output_json prediction_model={payload.get('prediction_model')} != current model={model}"
        )
    if int(payload.get("context_size") or 0) != context_size:
        raise ValueError(
            f"Existing output_json context_size={payload.get('context_size')} != current context_size={context_size}"
        )
    if bool(payload.get("include_current_turn")) != include_current_turn:
        raise ValueError(
            f"Existing output_json include_current_turn={payload.get('include_current_turn')} "
            f"!= current include_current_turn={include_current_turn}"
        )
    existing_relation_gold_source = clean_text(payload.get("relation_gold_source") or payload.get("relation_gold_glob"))
    if existing_relation_gold_source != relation_gold_source:
        raise ValueError(
            f"Existing output_json relation_gold_source={existing_relation_gold_source} "
            f"!= current relation_gold_source={relation_gold_source}"
        )
    if clean_text(payload.get("target_speaker_filter") or "therapist") != target_speaker_filter:
        raise ValueError(
            f"Existing output_json target_speaker_filter={payload.get('target_speaker_filter')} "
            f"!= current target_speaker_filter={target_speaker_filter}"
        )
    if clean_text(payload.get("group_mode") or "couple") != group_mode:
        raise ValueError(
            f"Existing output_json group_mode={payload.get('group_mode')} != current group_mode={group_mode}"
        )
    return list(payload.get("details", []) or [])


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Evaluate relation-cycle prediction with four multimodal modes."
    )
    parser.add_argument("--data-root", default="data")
    parser.add_argument("--output-json", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--sample-ids-file", default="")
    parser.add_argument("--limit-samples", type=int, default=0)
    parser.add_argument("--context-size", type=int, default=-1, help="Number of prior rows to include. Use 0 for none, -1 for all.")
    parser.add_argument("--include-current-turn", action="store_true", help="Also provide the current target turn to the model. Default is prior-context only.")
    parser.add_argument(
        "--target-speaker-filter",
        choices=sorted(TARGET_SPEAKER_FILTERS),
        default="therapist",
        help="Which gold-labeled rows to evaluate: therapist-only or any speaker.",
    )
    parser.add_argument("--mode", choices=sorted(VALID_TEXT_MODES | {"text"}), default="dialogue")
    parser.add_argument("--batch-size", type=int, default=3)
    parser.add_argument(
        "--relation-gold-glob",
        default="",
        help=(
            "Optional legacy external relation-gold glob. If omitted, relation_cycle_state "
            "annotations are read directly from rows in --data-root."
        ),
    )
    parser.add_argument("--group-mode", choices=["couple", "family"], default="couple")
    parser.add_argument(
        "--api-base",
        default=os.getenv("BAILIAN_API_BASE")
        or os.getenv("DASHSCOPE_API_BASE")
        or os.getenv("OPENAI_API_BASE")
        or os.getenv("OPENAI_BASE_URL")
        or DEFAULT_API_BASE,
    )
    parser.add_argument(
        "--api-key",
        default=os.getenv("BAILIAN_API_KEY") or os.getenv("DASHSCOPE_API_KEY") or os.getenv("OPENAI_API_KEY"),
    )
    parser.add_argument("--model", default=os.getenv("EVAL_MODEL") or os.getenv("BAILIAN_MODEL") or "deepseek-v4-flash")
    parser.add_argument("--max-tokens", type=int, default=1800)
    parser.add_argument("--disable-thinking", action="store_true")
    args = parser.parse_args()

    mode = MODE_ALIASES.get(args.mode, args.mode)
    if mode not in VALID_TEXT_MODES:
        raise ValueError(f"Unsupported mode: {args.mode}")
    if args.context_size < -1:
        raise ValueError("context_size must be >= -1. Use -1 for all prior rows.")
    if args.batch_size <= 0:
        raise ValueError("batch_size must be >= 1.")
    if args.target_speaker_filter not in TARGET_SPEAKER_FILTERS:
        raise ValueError(f"Unsupported target_speaker_filter: {args.target_speaker_filter}")
    if not args.api_key:
        raise ValueError("Missing API key.")
    if not args.model:
        raise ValueError("Missing model.")

    data_root = Path(args.data_root).expanduser().resolve()
    output_json = Path(args.output_json).expanduser().resolve()
    output_xlsx = Path(args.output_xlsx).expanduser().resolve()
    api_base = normalize_api_base(args.api_base)
    allowed_sample_ids = load_sample_id_filter(args.sample_ids_file)
    relation_labels = relation_labels_for_group_mode(args.group_mode)

    samples = load_data_samples(data_root, ROOT_DIR)
    if allowed_sample_ids:
        samples = [sample for sample in samples if sample["sample_id"] in allowed_sample_ids]
    if args.limit_samples > 0:
        samples = samples[: args.limit_samples]
    relation_gold_source = clean_text(args.relation_gold_glob) or "embedded"
    gold_map = (
        load_relation_gold_map(args.relation_gold_glob)
        if clean_text(args.relation_gold_glob)
        else load_embedded_relation_gold_map(samples)
    )

    detail_rows = load_existing_details(
        output_json,
        mode,
        args.model,
        args.context_size,
        relation_gold_source,
        args.group_mode,
        args.include_current_turn,
        args.target_speaker_filter,
    )
    completed_sample_ids = {row["sample_id"] for row in build_sample_rows_from_details(detail_rows)}

    targets_by_sample: Dict[str, List[Dict[str, Any]]] = {}
    total_checkpoint_count = len(detail_rows)
    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            continue

        targets = build_targets(sample, args.context_size, gold_map, args.target_speaker_filter)
        if not targets:
            continue
        gold_group_modes = {
            clean_text(gold_map[(sample_id, target["row_index"])].get("group_mode") or "couple")
            for target in targets
            if (sample_id, target["row_index"]) in gold_map
        }
        gold_group_modes.discard("")
        if gold_group_modes and gold_group_modes != {args.group_mode}:
            raise ValueError(
                f"Gold relation group_mode mismatch for {sample_id}: found {sorted(gold_group_modes)} but current group_mode={args.group_mode}"
            )
        targets_by_sample[sample_id] = targets
        total_checkpoint_count += len(targets)

    total_progress = tqdm(
        total=total_checkpoint_count,
        initial=len(detail_rows),
        desc="relation total",
        unit="turn",
        dynamic_ncols=True,
    )

    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            print(f"[skip] {sample_id}", flush=True)
            continue

        targets = targets_by_sample.get(sample_id, [])
        if not targets:
            continue

        target_by_turn = {target["turn_id"]: target for target in targets}
        pred_by_turn: Dict[str, Dict[str, Any]] = {}
        with tqdm(total=len(targets), desc=f"relation {sample_id}", unit="turn", dynamic_ncols=True, disable=True) as progress:
            for batch in batched(targets, args.batch_size):
                before = len(pred_by_turn)
                records = [relation_prompt_record(target, mode, args.context_size, args.include_current_turn, relation_labels) for target in batch]
                try:
                    payload = request_text_predictions(
                        api_base=api_base,
                        api_key=args.api_key,
                        model=args.model,
                        sample_id=sample_id,
                        group_mode=args.group_mode,
                        targets=records,
                        mode=mode,
                        context_size=args.context_size,
                        include_current_turn=args.include_current_turn,
                        relation_labels=relation_labels,
                        max_tokens=args.max_tokens,
                        disable_thinking=args.disable_thinking,
                    )
                    merge_relation_prediction_items(pred_by_turn, payload, batch)
                except Exception as exc:
                    print(
                        f"[warn] {sample_id}: relation batch request failed for {len(batch)} turns; "
                        f"retrying individually. error={clean_text(str(exc))[:300]}",
                        flush=True,
                    )
                    for target in batch:
                        single_before = len(pred_by_turn)
                        try:
                            payload = request_text_predictions(
                                api_base=api_base,
                                api_key=args.api_key,
                                model=args.model,
                                sample_id=sample_id,
                                group_mode=args.group_mode,
                                targets=[relation_prompt_record(target, mode, args.context_size, args.include_current_turn, relation_labels)],
                                mode=mode,
                                context_size=args.context_size,
                                include_current_turn=args.include_current_turn,
                                relation_labels=relation_labels,
                                max_tokens=min(900, args.max_tokens),
                                disable_thinking=args.disable_thinking,
                            )
                            raw_items = coerce_single_relation_prediction_items(payload, target["turn_id"])
                            for item in raw_items:
                                turn_id = clean_text(item.get("turn_id")) or target["turn_id"]
                                if turn_id == target["turn_id"]:
                                    pred_by_turn[turn_id] = item
                        except Exception as single_exc:
                            pred_by_turn[target["turn_id"]] = build_relation_fallback_prediction(
                                target,
                                relation_labels,
                                single_exc,
                            )
                            print(
                                f"[warn] {sample_id}: fallback relation_top3 for {target['turn_id']} after repeated "
                                f"request failure. error={clean_text(str(single_exc))[:240]}",
                                flush=True,
                            )
                        delta = max(0, len(pred_by_turn) - single_before)
                        progress.update(delta)
                        total_progress.update(delta)
                        progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                        total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                        time.sleep(0.3)
                    continue
                delta = max(0, len(pred_by_turn) - before)
                progress.update(delta)
                total_progress.update(delta)
                progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                time.sleep(0.3)

            missing = [target["turn_id"] for target in targets if target["turn_id"] not in pred_by_turn]
            if missing:
                missing_targets = [target for target in targets if target["turn_id"] in set(missing)]
                for target in missing_targets:
                    before = len(pred_by_turn)
                    try:
                        payload = request_text_predictions(
                            api_base=api_base,
                            api_key=args.api_key,
                            model=args.model,
                            sample_id=sample_id,
                            group_mode=args.group_mode,
                            targets=[relation_prompt_record(target, mode, args.context_size, args.include_current_turn, relation_labels)],
                            mode=mode,
                            context_size=args.context_size,
                            include_current_turn=args.include_current_turn,
                            relation_labels=relation_labels,
                            max_tokens=min(900, args.max_tokens),
                            disable_thinking=args.disable_thinking,
                        )
                        raw_items = coerce_single_relation_prediction_items(payload, target["turn_id"])
                        for item in raw_items:
                            turn_id = clean_text(item.get("turn_id")) or target["turn_id"]
                            if turn_id == target["turn_id"]:
                                pred_by_turn[turn_id] = item
                    except Exception as exc:
                        pred_by_turn[target["turn_id"]] = build_relation_fallback_prediction(
                            target,
                            relation_labels,
                            exc,
                        )
                        print(
                            f"[warn] {sample_id}: fallback relation_top3 for missing turn {target['turn_id']}. "
                            f"error={clean_text(str(exc))[:240]}",
                            flush=True,
                        )
                    delta = max(0, len(pred_by_turn) - before)
                    progress.update(delta)
                    total_progress.update(delta)
                    progress.set_postfix_str(f"pred={len(pred_by_turn)}/{len(targets)}")
                    total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_turn)}/{len(targets)}")
                    time.sleep(0.3)
                missing = [target["turn_id"] for target in targets if target["turn_id"] not in pred_by_turn]
        if missing:
            print(
                f"[warn] {sample_id}: filling {len(missing)} unresolved missing turns with fallback relation_top3.",
                flush=True,
            )
            for turn_id in missing:
                target = target_by_turn[turn_id]
                pred_by_turn[turn_id] = build_relation_fallback_prediction(
                    target,
                    relation_labels,
                    RuntimeError("missing prediction after batch and single-turn retries"),
                )

        sample_detail_rows: List[Dict[str, Any]] = []
        for target in targets:
            pred = pred_by_turn[target["turn_id"]]
            relation_top3 = normalize_relation_top3(pred.get("relation_top3", []), relation_labels)
            metrics = rank_metrics(target["gold_relation_state"], relation_top3)
            row = {
                "couple": target["couple"],
                "sample_id": sample_id,
                "group_mode": args.group_mode,
                "_relation_labels": relation_labels,
                "turn_id": target["turn_id"],
                "row_index": target["row_index"],
                "timestamp": target["timestamp"],
                "target_primary_speaker": target["target_primary_speaker"],
                "gold_relation_state": target["gold_relation_state"],
                "pred_relation_top1": relation_top3[0],
                "pred_relation_top2": relation_top3[1],
                "pred_relation_top3": relation_top3[2],
                "relation_gold_rank": metrics["gold_rank"],
                "relation_recall_at_1": metrics["recall_at_1"],
                "relation_recall_at_2": metrics["recall_at_2"],
                "relation_recall_at_3": metrics["recall_at_3"],
                "relation_mrr": metrics["mrr"],
                "mode": mode,
                "dialogue": target["dialogue"],
                "reason": clean_text(pred.get("reason")),
            }
            detail_rows.append(row)
            sample_detail_rows.append(row)

        sample_metrics = summarize_rows(sample_detail_rows)
        completed_sample_ids.add(sample_id)

        overall_rows = [{"scope": "overall", **summarize_rows(detail_rows)}]
        couple_rows = build_couple_rows_from_details(detail_rows)
        sample_rows = build_sample_rows_from_details(detail_rows)
        relation_confusion_rows = build_relation_confusion_rows(detail_rows)
        result_payload = build_result_payload(
            mode=mode,
            model=args.model,
            api_base=api_base,
            data_root=data_root,
            relation_gold_source=relation_gold_source,
            group_mode=args.group_mode,
            context_size=args.context_size,
            include_current_turn=args.include_current_turn,
            target_speaker_filter=args.target_speaker_filter,
            detail_rows=detail_rows,
            sample_rows=sample_rows,
            overall_rows=overall_rows,
            couple_rows=couple_rows,
            relation_confusion_rows=relation_confusion_rows,
        )
        output_json.parent.mkdir(parents=True, exist_ok=True)
        output_json.write_text(json.dumps(result_payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(
            f"[done] {sample_id}: turns={sample_metrics['turn_count']} "
            f"r1={sample_metrics['relation_recall_at_1']:.3f}",
            flush=True,
        )

    total_progress.close()

    overall_rows = [{"scope": "overall", **summarize_rows(detail_rows)}]
    couple_rows = build_couple_rows_from_details(detail_rows)
    sample_rows = build_sample_rows_from_details(detail_rows)
    relation_confusion_rows = build_relation_confusion_rows(detail_rows)
    result_payload = build_result_payload(
        mode=mode,
        model=args.model,
        api_base=api_base,
        data_root=data_root,
        relation_gold_source=relation_gold_source,
        group_mode=args.group_mode,
        context_size=args.context_size,
        include_current_turn=args.include_current_turn,
        target_speaker_filter=args.target_speaker_filter,
        detail_rows=detail_rows,
        sample_rows=sample_rows,
        overall_rows=overall_rows,
        couple_rows=couple_rows,
        relation_confusion_rows=relation_confusion_rows,
    )
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(result_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_result_workbook(
        output_xlsx,
        overall_rows,
        couple_rows,
        sample_rows,
        detail_rows,
        relation_confusion_rows,
        result_payload["top1_class_metrics"],
    )
    print(f"[written] {output_json}", flush=True)
    print(f"[written] {output_xlsx}", flush=True)


if __name__ == "__main__":
    main()
