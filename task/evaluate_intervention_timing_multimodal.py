import argparse
import json
import os
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
import requests
try:
    from tqdm.auto import tqdm
except Exception:  # pragma: no cover - tqdm is optional for script progress.
    def tqdm(iterable=None, **kwargs):
        return iterable if iterable is not None else _NullProgress()


    class _NullProgress:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def update(self, n=1):
            return None

        def set_postfix_str(self, *args, **kwargs):
            return None


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from task.eval_common import (
    DEFAULT_API_BASE,
    call_json_model,
    clean_text,
    load_data_samples,
    load_sample_id_filter,
    mask_sensitive_terms,
    normalize_api_base,
)
from task.multimodal_eval_common import (
    MODE_ALIASES,
    VALID_TEXT_MODES,
    build_prior_context_rows,
    load_review_manifest,
    normalize_background_dialogue,
    normalize_interaction_behavior,
    normalize_named_descriptions,
    prompt_record_for_mode,
    row_timestamp,
)
from utils.project_env import load_project_env


load_project_env(ROOT_DIR)


TASK_NAME = "predict_therapist_intervention_timing_proxy_from_prior_context_only"


def batched(seq: Sequence[Dict[str, Any]], size: int) -> List[List[Dict[str, Any]]]:
    if size <= 0:
        raise ValueError("batch_size must be >= 1.")
    return [list(seq[i : i + size]) for i in range(0, len(seq), size)]


def is_therapist(label: Any) -> bool:
    return clean_text(label) == "Therapist"


def build_row_record(
    row: Dict[str, Any],
    row_index: int,
    manifest_rows: Dict[int, Dict[str, Any]],
) -> Dict[str, Any]:
    clip_path = ""
    clip_duration_sec = 0.0
    manifest_item = manifest_rows.get(row_index) or {}
    clip_rel = clean_text(manifest_item.get("clip_path"))
    if clip_rel:
        resolved = (ROOT_DIR / clip_rel).resolve()
        if resolved.exists():
            clip_path = str(resolved)
    try:
        clip_duration_sec = float(manifest_item.get("duration_sec") or 0.0)
    except Exception:
        clip_duration_sec = 0.0

    return {
        "row_index": row_index,
        "timestamp": row_timestamp(row),
        "primary_speaker": clean_text(row.get("primary_speaker")),
        "target": [clean_text(item) for item in row.get("target") or [] if clean_text(item)],
        "dialogue_cleaned": clean_text(row.get("dialogue_cleaned")),
        "background_dialogue": normalize_background_dialogue(row.get("background_dialogue") or []),
        "tone_of_voice": clean_text(row.get("tone_of_voice")),
        "body_posture": normalize_named_descriptions(row.get("body_posture") or []),
        "facial_expressions": normalize_named_descriptions(row.get("facial_expressions") or []),
        "self_directed_behavior": normalize_named_descriptions(row.get("self_directed_behavior") or []),
        "interaction_behavior": normalize_interaction_behavior(row.get("interaction_behavior") or []),
        "clip_path": clip_path,
        "clip_duration_sec": max(0.0, clip_duration_sec),
    }


def even_sample(items: Sequence[Dict[str, Any]], limit: int) -> List[Dict[str, Any]]:
    if limit <= 0 or len(items) <= limit:
        return list(items)
    if limit == 1:
        return [items[len(items) // 2]]
    selected: List[Dict[str, Any]] = []
    for rank in range(limit):
        pos = round(rank * (len(items) - 1) / (limit - 1))
        selected.append(items[pos])
    seen = set()
    deduped: List[Dict[str, Any]] = []
    for item in selected:
        candidate_id = clean_text(item.get("candidate_id"))
        if candidate_id in seen:
            continue
        seen.add(candidate_id)
        deduped.append(item)
    if len(deduped) < limit:
        for item in items:
            candidate_id = clean_text(item.get("candidate_id"))
            if candidate_id in seen:
                continue
            seen.add(candidate_id)
            deduped.append(item)
            if len(deduped) >= limit:
                break
    return deduped[:limit]


def build_targets(
    sample: Dict[str, Any],
    context_size: int,
    max_candidates_per_class_per_sample: int,
) -> List[Dict[str, Any]]:
    rows = sample["rows"]
    manifest_rows = load_review_manifest(sample["sample_id"])
    positives: List[Dict[str, Any]] = []
    negatives: List[Dict[str, Any]] = []

    for row_index in range(2, len(rows) + 1):
        current_row = rows[row_index - 1]
        prev_row = rows[row_index - 2]
        if is_therapist(current_row.get("primary_speaker")) and not is_therapist(prev_row.get("primary_speaker")):
            positives.append(
                {
                    "sample_id": sample["sample_id"],
                    "couple": sample["couple"],
                    "candidate_id": f"{sample['sample_id']}__pos_after_row{row_index - 1}",
                    "candidate_type": "actual_therapist_entry",
                    "boundary_type": "therapist_entry",
                    "context_end_row_index": row_index - 1,
                    "context_end_timestamp": row_timestamp(prev_row),
                    "context_end_speaker": clean_text(prev_row.get("primary_speaker")),
                    "context_end_dialogue": clean_text(prev_row.get("dialogue_cleaned")),
                    "context_end_row_record": build_row_record(prev_row, row_index - 1, manifest_rows),
                    "actual_next_row_index": row_index,
                    "actual_next_timestamp": row_timestamp(current_row),
                    "actual_next_speaker": clean_text(current_row.get("primary_speaker")),
                    "actual_next_dialogue": clean_text(current_row.get("dialogue_cleaned")),
                    "gold_should_speak": "yes",
                    "prior_context": build_prior_context_rows(rows, manifest_rows, row_index, context_size),
                }
            )

    for row_index in range(1, len(rows)):
        current_row = rows[row_index - 1]
        next_row = rows[row_index]
        if is_therapist(current_row.get("primary_speaker")) or is_therapist(next_row.get("primary_speaker")):
            continue
        negatives.append(
            {
                "sample_id": sample["sample_id"],
                "couple": sample["couple"],
                "candidate_id": f"{sample['sample_id']}__neg_after_row{row_index}",
                "candidate_type": "continue_listening",
                "boundary_type": (
                    "same_speaker_continuation"
                    if clean_text(current_row.get("primary_speaker")) == clean_text(next_row.get("primary_speaker"))
                    else "non_therapist_exchange"
                ),
                "context_end_row_index": row_index,
                "context_end_timestamp": row_timestamp(current_row),
                "context_end_speaker": clean_text(current_row.get("primary_speaker")),
                "context_end_dialogue": clean_text(current_row.get("dialogue_cleaned")),
                "context_end_row_record": build_row_record(current_row, row_index, manifest_rows),
                "actual_next_row_index": row_index + 1,
                "actual_next_timestamp": row_timestamp(next_row),
                "actual_next_speaker": clean_text(next_row.get("primary_speaker")),
                "actual_next_dialogue": clean_text(next_row.get("dialogue_cleaned")),
                "gold_should_speak": "no",
                "prior_context": build_prior_context_rows(rows, manifest_rows, row_index + 1, context_size),
            }
        )

    if max_candidates_per_class_per_sample > 0:
        positives = even_sample(positives, max_candidates_per_class_per_sample)
        negatives = even_sample(negatives, max_candidates_per_class_per_sample)

    balanced_count = min(len(positives), len(negatives))
    if balanced_count <= 0:
        return []

    positives = even_sample(positives, balanced_count)
    negatives = even_sample(negatives, balanced_count)
    return sorted(positives + negatives, key=lambda item: item["context_end_row_index"])


def timing_prompt_record(target: Dict[str, Any], mode: str, context_size: int) -> Dict[str, Any]:
    record: Dict[str, Any] = {
        "candidate_id": target["candidate_id"],
        "candidate_type": target["candidate_type"],
        "boundary_type": target["boundary_type"],
        "decision_after_row": prompt_record_for_mode(target["context_end_row_record"], mode),
    }
    if context_size != 0:
        record["prior_context"] = [prompt_record_for_mode(item, mode) for item in target["prior_context"]]
    else:
        record["prior_context"] = []
    return record


def sanitize_timing_records(records: Sequence[Dict[str, Any]], level: str) -> List[Dict[str, Any]]:
    sanitized: List[Dict[str, Any]] = []
    for record in records:
        item = json.loads(json.dumps(record, ensure_ascii=False))
        if level == "none":
            sanitized.append(item)
            continue

        containers = list(item.get("prior_context", [])) + [item.get("decision_after_row", {})]
        for row in containers:
            if not isinstance(row, dict):
                continue
            if "dialogue_cleaned" in row:
                if level == "mask":
                    row["dialogue_cleaned"] = mask_sensitive_terms(clean_text(row.get("dialogue_cleaned")))
                else:
                    row["dialogue_cleaned"] = "[redacted sensitive prior dialogue]"
            for bg in row.get("background_dialogue", []) or []:
                if "content" in bg:
                    if level == "mask":
                        bg["content"] = mask_sensitive_terms(clean_text(bg.get("content")))
                    else:
                        bg["content"] = "[redacted sensitive background dialogue]"
        sanitized.append(item)
    return sanitized


def normalize_group_modes(values: Sequence[Any]) -> List[str]:
    seen = set()
    normalized: List[str] = []
    for value in values:
        label = clean_text(value) or "couple"
        if label in seen:
            continue
        seen.add(label)
        normalized.append(label)
    return normalized or ["couple"]


def prediction_prompt(
    sample_id: str,
    group_mode: str,
    mode: str,
    records: Sequence[Dict[str, Any]],
    context_size: int,
) -> str:
    session_phrase = "family-therapy interaction" if group_mode == "family" else "couples-therapy interaction"
    no_guidance = (
        "Prefer `no` when a family member is still unfolding a thought, vulnerable material is still emerging, or the family exchange is still productively developing without therapist interruption.\n"
        if group_mode == "family"
        else "Prefer `no` when a speaker is still unfolding a thought, vulnerable material is still emerging, or the couple's exchange is still productively developing without therapist interruption.\n"
    )
    mode_guidance = {
        "dialogue": "Use only language-side evidence from prior_context and decision_after_row: primary_speaker, target, dialogue_cleaned, and background_dialogue.",
        "audio": "Use language-side evidence together with tone_of_voice. Do not use visual annotations.",
        "visual": "Use language plus visual annotation evidence from body_posture, facial_expressions, self_directed_behavior, and interaction_behavior.",
        "audiovisual": "Use language, visual annotations, and tone_of_voice.",
    }
    context_guidance = (
        "No earlier prior_context rows are provided beyond the decision_after_row.\n"
        if context_size == 0
        else (
            "Use all available earlier rows in prior_context before the decision point.\n"
            if context_size < 0
            else f"Use at most the previous {context_size} rows in prior_context before the decision point.\n"
        )
    )
    return (
        f"You are evaluating therapist intervention timing in a {session_phrase}.\n"
        "For each candidate decision point, decide whether the therapist should take the NEXT turn immediately after decision_after_row.\n"
        "Return `yes` if a therapist intervention is timely and clinically helpful NOW.\n"
        "Return `no` if the therapist should keep listening and let the clients continue.\n"
        f"{context_guidance}"
        f"Mode: {mode}\n"
        f"{mode_guidance[mode]}\n"
        "Use a concise clinical standard.\n"
        "Prefer `yes` when the process has crystallized enough that a question, reflection, validation, containment, reframing, or process intervention would likely help now.\n"
        f"{no_guidance}"
        "Judge only from the provided context. Do not assume access to the unseen next turn.\n"
        "Echo the exact Sample value and exact input candidate_id values from Targets.\n"
        "Output JSON only with this schema:\n"
        "{"
        "\"sample_id\":\"<echo the exact Sample value below>\","
        "\"predictions\":["
        "{"
        "\"candidate_id\":\"<echo one exact candidate_id from Targets>\","
        "\"should_speak\":\"yes\","
        "\"confidence\":0.74,"
        "\"reason\":\"short English reason based only on the provided context\""
        "}"
        "]"
        "}\n"
        "No markdown. JSON only.\n"
        f"Sample: {sample_id}\n"
        f"Targets: {json.dumps(list(records), ensure_ascii=False)}"
    )


def request_text_predictions(
    *,
    api_base: str,
    api_key: str,
    model: str,
    sample_id: str,
    group_mode: str,
    targets: Sequence[Dict[str, Any]],
    mode: str,
    context_size: int,
    max_tokens: int,
    disable_thinking: bool,
) -> Dict[str, Any]:
    prompts = [
        prediction_prompt(sample_id, group_mode, mode, sanitize_timing_records(targets, "none"), context_size),
        prediction_prompt(sample_id, group_mode, mode, sanitize_timing_records(targets, "mask"), context_size),
        prediction_prompt(sample_id, group_mode, mode, sanitize_timing_records(targets, "redact"), context_size),
    ]
    last_error: Optional[Exception] = None
    for prompt_index, prompt in enumerate(prompts):
        for attempt in range(1, 4):
            try:
                return call_json_model(
                    api_base=api_base,
                    api_key=api_key,
                    model=model,
                    prompt=prompt,
                    max_tokens=max_tokens,
                    disable_thinking=disable_thinking,
                )
            except requests.HTTPError as exc:
                last_error = exc
                if "data_inspection_failed" in str(exc) and prompt_index < len(prompts) - 1:
                    break
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
            except json.JSONDecodeError as exc:
                last_error = exc
                if attempt == 3:
                    break
                time.sleep(2 * attempt)
            except (requests.RequestException, OSError) as exc:
                last_error = exc
                if attempt == 3:
                    raise
                time.sleep(3 * attempt)
    if last_error is not None:
        raise last_error
    raise RuntimeError("Unknown error while requesting timing predictions")


def extract_prediction_items(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    items = payload.get("predictions", [])
    if isinstance(items, dict):
        items = [items]
    if not isinstance(items, list):
        items = []
    items = [item for item in items if isinstance(item, dict)]
    if not items and isinstance(payload, dict) and payload.get("should_speak"):
        items = [payload]
    return items


def merge_prediction_items(
    pred_by_id: Dict[str, Dict[str, Any]],
    payload: Dict[str, Any],
    batch: Sequence[Dict[str, Any]],
) -> None:
    for item in extract_prediction_items(payload):
        candidate_id = clean_text(item.get("candidate_id"))
        if not candidate_id and len(batch) == 1:
            candidate_id = batch[0]["candidate_id"]
        if candidate_id:
            pred_by_id[candidate_id] = item


def build_fallback_prediction(target: Dict[str, Any], error: Exception) -> Dict[str, Any]:
    reason = clean_text(str(error)).replace("\n", " ")
    if len(reason) > 240:
        reason = reason[:237] + "..."
    return {
        "candidate_id": target["candidate_id"],
        "should_speak": "no",
        "confidence": 0.0,
        "reason": f"fallback_after_error: {reason}",
    }


def normalize_yes_no(value: Any) -> str:
    if isinstance(value, bool):
        return "yes" if value else "no"
    text = clean_text(value).lower()
    if text in {"yes", "y", "true", "1", "speak", "should_speak"}:
        return "yes"
    if text in {"no", "n", "false", "0", "listen", "keep_listening", "do_not_speak"}:
        return "no"
    if "yes" in text:
        return "yes"
    if "no" in text:
        return "no"
    return ""


def parse_confidence(value: Any) -> float | str:
    if value in (None, ""):
        return ""
    try:
        score = float(value)
    except Exception:
        return clean_text(value)
    return max(0.0, min(1.0, score))


def classification_metrics(rows: Sequence[Dict[str, Any]]) -> Dict[str, float]:
    count = len(rows)
    if count == 0:
        return {
            "candidate_count": 0,
            "accuracy": 0.0,
            "precision_yes": 0.0,
            "recall_yes": 0.0,
            "f1_yes": 0.0,
            "precision_no": 0.0,
            "recall_no": 0.0,
            "f1_no": 0.0,
            "balanced_accuracy": 0.0,
            "tp": 0,
            "tn": 0,
            "fp": 0,
            "fn": 0,
        }

    tp = sum(1 for row in rows if row["gold_should_speak"] == "yes" and row["pred_should_speak"] == "yes")
    tn = sum(1 for row in rows if row["gold_should_speak"] == "no" and row["pred_should_speak"] == "no")
    fp = sum(1 for row in rows if row["gold_should_speak"] == "no" and row["pred_should_speak"] == "yes")
    fn = sum(1 for row in rows if row["gold_should_speak"] == "yes" and row["pred_should_speak"] == "no")

    def safe_div(numerator: float, denominator: float) -> float:
        return numerator / denominator if denominator else 0.0

    precision_yes = safe_div(tp, tp + fp)
    recall_yes = safe_div(tp, tp + fn)
    f1_yes = safe_div(2 * precision_yes * recall_yes, precision_yes + recall_yes)
    precision_no = safe_div(tn, tn + fn)
    recall_no = safe_div(tn, tn + fp)
    f1_no = safe_div(2 * precision_no * recall_no, precision_no + recall_no)
    accuracy = safe_div(tp + tn, count)
    balanced_accuracy = (recall_yes + recall_no) / 2 if count else 0.0

    return {
        "candidate_count": count,
        "accuracy": accuracy,
        "precision_yes": precision_yes,
        "recall_yes": recall_yes,
        "f1_yes": f1_yes,
        "precision_no": precision_no,
        "recall_no": recall_no,
        "f1_no": f1_no,
        "balanced_accuracy": balanced_accuracy,
        "tp": tp,
        "tn": tn,
        "fp": fp,
        "fn": fn,
    }


def build_sample_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_sample: Dict[tuple[str, str, str], List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_sample[(row["couple"], row["sample_id"], clean_text(row.get("group_mode") or "couple"))].append(row)
    return [
        {"couple": couple, "sample_id": sample_id, "group_mode": group_mode, **classification_metrics(rows)}
        for (couple, sample_id, group_mode), rows in sorted(by_sample.items(), key=lambda item: item[0][1])
    ]


def build_couple_rows_from_details(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_couple: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_couple[row["couple"]].append(row)
    result: List[Dict[str, Any]] = []
    for couple, rows in sorted(by_couple.items(), key=lambda item: item[0]):
        group_modes = normalize_group_modes([row.get("group_mode") for row in rows])
        result.append(
            {
                "scope": couple,
                "group_mode": group_modes[0] if len(group_modes) == 1 else "mixed",
                **classification_metrics(rows),
            }
        )
    return result


def build_confusion_rows(detail_rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    counter: Counter = Counter((row["gold_should_speak"], row["pred_should_speak"]) for row in detail_rows)
    return [
        {"gold_should_speak": gold, "pred_should_speak": pred, "count": count}
        for (gold, pred), count in counter.most_common()
    ]


def write_result_workbook(
    output_path: Path,
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    detail_rows: Sequence[Dict[str, Any]],
    confusion_rows: Sequence[Dict[str, Any]],
) -> None:
    workbook = Workbook()
    ws_overall = workbook.active
    ws_overall.title = "overall"
    ws_overall.append(
        [
            "scope",
            "candidate_count",
            "accuracy",
            "precision_yes",
            "recall_yes",
            "f1_yes",
            "precision_no",
            "recall_no",
            "f1_no",
            "balanced_accuracy",
            "tp",
            "tn",
            "fp",
            "fn",
        ]
    )
    for row in overall_rows:
        ws_overall.append(
            [
                row["scope"],
                row["candidate_count"],
                row["accuracy"],
                row["precision_yes"],
                row["recall_yes"],
                row["f1_yes"],
                row["precision_no"],
                row["recall_no"],
                row["f1_no"],
                row["balanced_accuracy"],
                row["tp"],
                row["tn"],
                row["fp"],
                row["fn"],
            ]
        )

    ws_couples = workbook.create_sheet("couples")
    ws_couples.append(
        [
            "scope",
            "group_mode",
            "candidate_count",
            "accuracy",
            "precision_yes",
            "recall_yes",
            "f1_yes",
            "precision_no",
            "recall_no",
            "f1_no",
            "balanced_accuracy",
            "tp",
            "tn",
            "fp",
            "fn",
        ]
    )
    for row in couple_rows:
        ws_couples.append(
            [
                row["scope"],
                row["group_mode"],
                row["candidate_count"],
                row["accuracy"],
                row["precision_yes"],
                row["recall_yes"],
                row["f1_yes"],
                row["precision_no"],
                row["recall_no"],
                row["f1_no"],
                row["balanced_accuracy"],
                row["tp"],
                row["tn"],
                row["fp"],
                row["fn"],
            ]
        )

    ws_samples = workbook.create_sheet("samples")
    ws_samples.append(
        [
            "couple",
            "sample_id",
            "group_mode",
            "candidate_count",
            "accuracy",
            "precision_yes",
            "recall_yes",
            "f1_yes",
            "precision_no",
            "recall_no",
            "f1_no",
            "balanced_accuracy",
            "tp",
            "tn",
            "fp",
            "fn",
        ]
    )
    for row in sample_rows:
        ws_samples.append(
            [
                row["couple"],
                row["sample_id"],
                row["group_mode"],
                row["candidate_count"],
                row["accuracy"],
                row["precision_yes"],
                row["recall_yes"],
                row["f1_yes"],
                row["precision_no"],
                row["recall_no"],
                row["f1_no"],
                row["balanced_accuracy"],
                row["tp"],
                row["tn"],
                row["fp"],
                row["fn"],
            ]
        )

    ws_details = workbook.create_sheet("details")
    ws_details.append(
        [
            "couple",
            "sample_id",
            "group_mode",
            "candidate_id",
            "candidate_type",
            "boundary_type",
            "context_end_row_index",
            "context_end_timestamp",
            "context_end_speaker",
            "context_end_dialogue",
            "gold_should_speak",
            "pred_should_speak",
            "correct",
            "confidence",
            "mode",
            "reason",
            "actual_next_row_index",
            "actual_next_timestamp",
            "actual_next_speaker",
            "actual_next_dialogue",
        ]
    )
    for row in detail_rows:
        ws_details.append(
            [
                row["couple"],
                row["sample_id"],
                row["group_mode"],
                row["candidate_id"],
                row["candidate_type"],
                row["boundary_type"],
                row["context_end_row_index"],
                row["context_end_timestamp"],
                row["context_end_speaker"],
                row["context_end_dialogue"],
                row["gold_should_speak"],
                row["pred_should_speak"],
                row["correct"],
                row["confidence"],
                row["mode"],
                row["reason"],
                row["actual_next_row_index"],
                row["actual_next_timestamp"],
                row["actual_next_speaker"],
                row["actual_next_dialogue"],
            ]
        )

    ws_conf = workbook.create_sheet("confusions")
    ws_conf.append(["gold_should_speak", "pred_should_speak", "count"])
    for row in confusion_rows:
        ws_conf.append([row["gold_should_speak"], row["pred_should_speak"], row["count"]])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_result_payload(
    *,
    group_modes: Sequence[str],
    mode: str,
    model: str,
    api_base: str,
    data_root: Path,
    context_size: int,
    max_candidates_per_class_per_sample: int,
    detail_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    overall_rows: Sequence[Dict[str, Any]],
    couple_rows: Sequence[Dict[str, Any]],
    confusion_rows: Sequence[Dict[str, Any]],
) -> Dict[str, Any]:
    normalized_group_modes = normalize_group_modes(group_modes or [row.get("group_mode") for row in detail_rows])
    return {
        "task": TASK_NAME,
        "group_mode": normalized_group_modes[0] if len(normalized_group_modes) == 1 else "mixed",
        "group_modes": normalized_group_modes,
        "mode": mode,
        "prediction_model": model,
        "prediction_api_base": api_base,
        "data_root": str(data_root),
        "context_size": context_size,
        "max_candidates_per_class_per_sample": max_candidates_per_class_per_sample,
        "gold_policy": "Observational proxy: yes means the actual next row is a therapist entry after a non-therapist row; no means the actual next row remains non-therapist, so continued listening occurred.",
        "balance_policy": "Per-sample strict balance between positive therapist-entry candidates and negative continue-listening candidates.",
        "overall": list(overall_rows),
        "couples": list(couple_rows),
        "samples": list(sample_rows),
        "details": list(detail_rows),
        "confusions": list(confusion_rows),
    }


def load_existing_details(
    output_json: Path,
    mode: str,
    model: str,
    context_size: int,
    max_candidates_per_class_per_sample: int,
    group_modes: Sequence[str],
) -> List[Dict[str, Any]]:
    if not output_json.exists():
        return []
    payload = json.loads(output_json.read_text(encoding="utf-8"))
    if clean_text(payload.get("task")) != TASK_NAME:
        return []
    if clean_text(payload.get("mode")) != mode:
        raise ValueError(f"Existing output_json mode={payload.get('mode')} != current mode={mode}")
    if clean_text(payload.get("prediction_model")) != model:
        raise ValueError(
            f"Existing output_json prediction_model={payload.get('prediction_model')} != current model={model}"
        )
    if int(payload.get("context_size") or 0) != context_size:
        raise ValueError(
            f"Existing output_json context_size={payload.get('context_size')} != current context_size={context_size}"
        )
    if int(payload.get("max_candidates_per_class_per_sample") or 0) != max_candidates_per_class_per_sample:
        raise ValueError(
            "Existing output_json max_candidates_per_class_per_sample="
            f"{payload.get('max_candidates_per_class_per_sample')} != current max_candidates_per_class_per_sample={max_candidates_per_class_per_sample}"
        )
    payload_group_modes = normalize_group_modes(payload.get("group_modes") or [payload.get("group_mode") or "couple"])
    expected_group_modes = normalize_group_modes(group_modes)
    if payload_group_modes != expected_group_modes:
        raise ValueError(
            f"Existing output_json group_modes={payload_group_modes} != current group_modes={expected_group_modes}"
        )
    return list(payload.get("details", []) or [])


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Evaluate therapist intervention timing from prior context with four multimodal modes."
    )
    parser.add_argument("--data-root", default="data")
    parser.add_argument("--output-json", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--sample-ids-file", default="")
    parser.add_argument("--limit-samples", type=int, default=0)
    parser.add_argument("--context-size", type=int, default=3, help="Number of prior rows to include. Use 0 for none, -1 for all.")
    parser.add_argument("--mode", choices=sorted(VALID_TEXT_MODES | {"text"}), default="dialogue")
    parser.add_argument("--batch-size", type=int, default=4)
    parser.add_argument("--max-candidates-per-class-per-sample", type=int, default=0)
    parser.add_argument(
        "--api-base",
        default=os.getenv("BAILIAN_API_BASE")
        or os.getenv("DASHSCOPE_API_BASE")
        or os.getenv("OPENAI_API_BASE")
        or os.getenv("OPENAI_BASE_URL")
        or DEFAULT_API_BASE,
    )
    parser.add_argument(
        "--api-key",
        default=os.getenv("BAILIAN_API_KEY") or os.getenv("DASHSCOPE_API_KEY") or os.getenv("OPENAI_API_KEY"),
    )
    parser.add_argument("--model", default=os.getenv("EVAL_MODEL") or os.getenv("BAILIAN_MODEL") or "deepseek-v4-flash")
    parser.add_argument("--max-tokens", type=int, default=1600)
    parser.add_argument("--disable-thinking", action="store_true")
    args = parser.parse_args()

    mode = MODE_ALIASES.get(args.mode, args.mode)
    if mode not in VALID_TEXT_MODES:
        raise ValueError(f"Unsupported mode: {args.mode}")
    if args.context_size < -1:
        raise ValueError("context_size must be >= -1. Use -1 for all prior rows.")
    if args.batch_size <= 0:
        raise ValueError("batch_size must be >= 1.")
    if not args.api_key:
        raise ValueError("Missing API key.")
    if not args.model:
        raise ValueError("Missing model.")

    data_root = Path(args.data_root).expanduser().resolve()
    output_json = Path(args.output_json).expanduser().resolve()
    output_xlsx = Path(args.output_xlsx).expanduser().resolve()
    api_base = normalize_api_base(args.api_base)
    allowed_sample_ids = load_sample_id_filter(args.sample_ids_file)

    samples = load_data_samples(data_root, ROOT_DIR)
    if allowed_sample_ids:
        samples = [sample for sample in samples if sample["sample_id"] in allowed_sample_ids]
    if args.limit_samples > 0:
        samples = samples[: args.limit_samples]
    selected_group_modes = normalize_group_modes(clean_text(sample.get("group_mode") or "couple") for sample in samples)

    detail_rows = load_existing_details(
        output_json,
        mode,
        args.model,
        args.context_size,
        args.max_candidates_per_class_per_sample,
        selected_group_modes,
    )
    completed_sample_ids = {row["sample_id"] for row in build_sample_rows_from_details(detail_rows)}

    targets_by_sample: Dict[str, List[Dict[str, Any]]] = {}
    total_checkpoint_count = len(detail_rows)
    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            continue
        targets = build_targets(sample, args.context_size, args.max_candidates_per_class_per_sample)
        if targets:
            targets_by_sample[sample_id] = targets
            total_checkpoint_count += len(targets)

    total_progress = tqdm(
        total=total_checkpoint_count,
        initial=len(detail_rows),
        desc="timing total",
        unit="candidate",
        dynamic_ncols=True,
    )

    for sample in samples:
        sample_id = sample["sample_id"]
        if sample_id in completed_sample_ids:
            print(f"[skip] {sample_id}", flush=True)
            continue

        targets = targets_by_sample.get(sample_id, [])
        if not targets:
            continue

        target_by_id = {target["candidate_id"]: target for target in targets}
        pred_by_id: Dict[str, Dict[str, Any]] = {}
        with tqdm(total=len(targets), desc=f"timing {sample_id}", unit="candidate", dynamic_ncols=True, disable=True) as progress:
            for batch in batched(targets, args.batch_size):
                before = len(pred_by_id)
                records = [timing_prompt_record(target, mode, args.context_size) for target in batch]
                try:
                    payload = request_text_predictions(
                        api_base=api_base,
                        api_key=args.api_key,
                        model=args.model,
                        sample_id=sample_id,
                        group_mode=clean_text(sample.get("group_mode") or "couple"),
                        targets=records,
                        mode=mode,
                        context_size=args.context_size,
                        max_tokens=args.max_tokens,
                        disable_thinking=args.disable_thinking,
                    )
                    merge_prediction_items(pred_by_id, payload, batch)
                except Exception as exc:
                    print(
                        f"[warn] {sample_id}: batch request failed for {len(batch)} candidates; "
                        f"retrying individually. error={clean_text(str(exc))[:300]}",
                        flush=True,
                    )
                    for target in batch:
                        single_before = len(pred_by_id)
                        try:
                            payload = request_text_predictions(
                                api_base=api_base,
                                api_key=args.api_key,
                                model=args.model,
                                sample_id=sample_id,
                                group_mode=clean_text(sample.get("group_mode") or "couple"),
                                targets=[timing_prompt_record(target, mode, args.context_size)],
                                mode=mode,
                                context_size=args.context_size,
                                max_tokens=min(900, args.max_tokens),
                                disable_thinking=args.disable_thinking,
                            )
                            merge_prediction_items(
                                pred_by_id,
                                payload,
                                [timing_prompt_record(target, mode, args.context_size)],
                            )
                        except Exception as single_exc:
                            pred_by_id[target["candidate_id"]] = build_fallback_prediction(target, single_exc)
                            print(
                                f"[warn] {sample_id}: fallback `no` for {target['candidate_id']} after repeated "
                                f"request failure. error={clean_text(str(single_exc))[:240]}",
                                flush=True,
                            )
                        delta = max(0, len(pred_by_id) - single_before)
                        progress.update(delta)
                        total_progress.update(delta)
                        progress.set_postfix_str(f"pred={len(pred_by_id)}/{len(targets)}")
                        total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_id)}/{len(targets)}")
                        time.sleep(0.3)
                    continue
                delta = max(0, len(pred_by_id) - before)
                progress.update(delta)
                total_progress.update(delta)
                progress.set_postfix_str(f"pred={len(pred_by_id)}/{len(targets)}")
                total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_id)}/{len(targets)}")
                time.sleep(0.3)

            missing = [target["candidate_id"] for target in targets if target["candidate_id"] not in pred_by_id]
            if missing:
                missing_targets = [target for target in targets if target["candidate_id"] in set(missing)]
                for target in missing_targets:
                    before = len(pred_by_id)
                    try:
                        payload = request_text_predictions(
                            api_base=api_base,
                            api_key=args.api_key,
                            model=args.model,
                            sample_id=sample_id,
                            group_mode=clean_text(sample.get("group_mode") or "couple"),
                            targets=[timing_prompt_record(target, mode, args.context_size)],
                            mode=mode,
                            context_size=args.context_size,
                            max_tokens=min(900, args.max_tokens),
                            disable_thinking=args.disable_thinking,
                        )
                        merge_prediction_items(
                            pred_by_id,
                            payload,
                            [timing_prompt_record(target, mode, args.context_size)],
                        )
                    except Exception as exc:
                        pred_by_id[target["candidate_id"]] = build_fallback_prediction(target, exc)
                        print(
                            f"[warn] {sample_id}: fallback `no` for missing candidate {target['candidate_id']}. "
                            f"error={clean_text(str(exc))[:240]}",
                            flush=True,
                        )
                    delta = max(0, len(pred_by_id) - before)
                    progress.update(delta)
                    total_progress.update(delta)
                    progress.set_postfix_str(f"pred={len(pred_by_id)}/{len(targets)}")
                    total_progress.set_postfix_str(f"sample={sample_id} pred={len(pred_by_id)}/{len(targets)}")
                    time.sleep(0.3)
                missing = [target["candidate_id"] for target in targets if target["candidate_id"] not in pred_by_id]
        if missing:
            print(
                f"[warn] {sample_id}: filling {len(missing)} unresolved missing candidates with fallback `no`.",
                flush=True,
            )
            for candidate_id in missing:
                target = target_by_id[candidate_id]
                pred_by_id[candidate_id] = build_fallback_prediction(
                    target,
                    RuntimeError("missing prediction after batch and single-candidate retries"),
                )

        sample_detail_rows: List[Dict[str, Any]] = []
        for target in targets:
            pred = pred_by_id[target["candidate_id"]]
            pred_should_speak = normalize_yes_no(pred.get("should_speak"))
            if pred_should_speak not in {"yes", "no"}:
                pred_should_speak = "no"
            row = {
                "couple": target["couple"],
                "sample_id": sample_id,
                "group_mode": clean_text(sample.get("group_mode") or "couple"),
                "candidate_id": target["candidate_id"],
                "candidate_type": target["candidate_type"],
                "boundary_type": target["boundary_type"],
                "context_end_row_index": target["context_end_row_index"],
                "context_end_timestamp": target["context_end_timestamp"],
                "context_end_speaker": target["context_end_speaker"],
                "context_end_dialogue": target["context_end_dialogue"],
                "gold_should_speak": target["gold_should_speak"],
                "pred_should_speak": pred_should_speak,
                "correct": 1.0 if pred_should_speak == target["gold_should_speak"] else 0.0,
                "confidence": parse_confidence(pred.get("confidence")),
                "mode": mode,
                "reason": clean_text(pred.get("reason")),
                "actual_next_row_index": target["actual_next_row_index"],
                "actual_next_timestamp": target["actual_next_timestamp"],
                "actual_next_speaker": target["actual_next_speaker"],
                "actual_next_dialogue": target["actual_next_dialogue"],
            }
            detail_rows.append(row)
            sample_detail_rows.append(row)

        sample_metrics = classification_metrics(sample_detail_rows)
        completed_sample_ids.add(sample_id)

        sample_rows = build_sample_rows_from_details(detail_rows)
        overall_rows = [{"scope": "overall", **classification_metrics(detail_rows)}]
        couple_rows = build_couple_rows_from_details(detail_rows)
        confusion_rows = build_confusion_rows(detail_rows)

        output_json.parent.mkdir(parents=True, exist_ok=True)
        output_json.write_text(
            json.dumps(
                build_result_payload(
                    group_modes=selected_group_modes,
                    mode=mode,
                    model=args.model,
                    api_base=api_base,
                    data_root=data_root,
                    context_size=args.context_size,
                    max_candidates_per_class_per_sample=args.max_candidates_per_class_per_sample,
                    detail_rows=detail_rows,
                    sample_rows=sample_rows,
                    overall_rows=overall_rows,
                    couple_rows=couple_rows,
                    confusion_rows=confusion_rows,
                ),
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        print(
            f"[done] {sample_id}: candidates={sample_metrics['candidate_count']} "
            f"acc={sample_metrics['accuracy']:.3f} "
            f"bal_acc={sample_metrics['balanced_accuracy']:.3f}",
            flush=True,
        )

    total_progress.close()

    overall_rows = [{"scope": "overall", **classification_metrics(detail_rows)}]
    sample_rows = build_sample_rows_from_details(detail_rows)
    couple_rows = build_couple_rows_from_details(detail_rows)
    confusion_rows = build_confusion_rows(detail_rows)

    result = build_result_payload(
        group_modes=selected_group_modes,
        mode=mode,
        model=args.model,
        api_base=api_base,
        data_root=data_root,
        context_size=args.context_size,
        max_candidates_per_class_per_sample=args.max_candidates_per_class_per_sample,
        detail_rows=detail_rows,
        sample_rows=sample_rows,
        overall_rows=overall_rows,
        couple_rows=couple_rows,
        confusion_rows=confusion_rows,
    )
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    write_result_workbook(output_xlsx, overall_rows, couple_rows, sample_rows, detail_rows, confusion_rows)
    print(f"[written] {output_json}", flush=True)
    print(f"[written] {output_xlsx}", flush=True)


if __name__ == "__main__":
    main()
