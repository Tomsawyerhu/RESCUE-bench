import argparse
import json
import os
import re
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Sequence

import requests
from openpyxl import Workbook


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))
from utils.project_env import load_project_env
from task.eval_common import infer_group_mode

load_project_env(ROOT_DIR)


DEFAULT_API_BASE = "https://dashscope.aliyuncs.com/compatible-mode/v1"

SENSITIVE_WORD_PATTERNS = [
    re.compile(r"\bfuck(?:ing|ed|er|ers)?\b", re.IGNORECASE),
    re.compile(r"\bshit(?:ty)?\b", re.IGNORECASE),
    re.compile(r"\bbitch(?:es)?\b", re.IGNORECASE),
    re.compile(r"\basshole(?:s)?\b", re.IGNORECASE),
    re.compile(r"\bdick\b", re.IGNORECASE),
    re.compile(r"\bpussy\b", re.IGNORECASE),
    re.compile(r"\bsex(?:ual)?\b", re.IGNORECASE),
    re.compile(r"\bpenis\b", re.IGNORECASE),
    re.compile(r"\bvagina\b", re.IGNORECASE),
    re.compile(r"\brape(?:d|s)?\b", re.IGNORECASE),
    re.compile(r"\bporn(?:ography|ographic)?\b", re.IGNORECASE),
]


def normalize_api_base(api_base: str) -> str:
    base = (api_base or "").strip().rstrip("/")
    if not base:
        raise ValueError("Missing API base.")
    return base if base.endswith("/v1") else f"{base}/v1"


def clean_text(value: Any) -> str:
    return str(value or "").strip()


def load_sample_id_filter(path: str | None) -> set[str]:
    if not path:
        return set()
    file_path = Path(path).expanduser().resolve()
    text = file_path.read_text(encoding="utf-8").strip()
    if not text:
        return set()
    try:
        payload = json.loads(text)
    except Exception:
        payload = None
    if isinstance(payload, list):
        return {clean_text(item) for item in payload if clean_text(item)}
    if isinstance(payload, dict) and isinstance(payload.get("sample_ids"), list):
        return {clean_text(item) for item in payload["sample_ids"] if clean_text(item)}
    return {clean_text(line) for line in text.splitlines() if clean_text(line)}


def mask_sensitive_terms(text: str) -> str:
    masked = text
    for pattern in SENSITIVE_WORD_PATTERNS:
        masked = pattern.sub(lambda match: match.group(0)[0] + "*" * max(1, len(match.group(0)) - 1), masked)
    return masked


def load_gold_dataset(path: Path) -> Dict[str, Dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    samples: Dict[str, Dict[str, Any]] = {}
    for sample in payload.get("samples", []):
        samples[sample["sample_id"]] = sample
    return samples


def batched(seq: Sequence[Dict[str, Any]], size: int) -> List[List[Dict[str, Any]]]:
    return [list(seq[i : i + size]) for i in range(0, len(seq), size)]


def build_targets(sample: Dict[str, Any]) -> List[Dict[str, Any]]:
    targets: List[Dict[str, Any]] = []
    for candidate in sorted(sample.get("candidates", []), key=lambda item: item["context_end_excel_row"]):
        targets.append(
            {
                "sample_id": sample["sample_id"],
                "candidate_id": candidate["candidate_id"],
                "sheet_name": candidate["sheet_name"],
                "episode_id": candidate["episode_id"],
                "segment_id": candidate["segment_id"],
                "context_end_excel_row": candidate["context_end_excel_row"],
                "context_end_timestamp": candidate["context_end_timestamp"],
                "prior_context": [
                    {
                        "excel_row": row["excel_row"],
                        "timestamp": row["timestamp"],
                        "speaker": row["speaker"],
                        "dialogue": row["dialogue"],
                    }
                    for row in candidate.get("prior_context", [])
                ],
                "gold_should_speak": candidate["gold_should_speak"],
                "context_end_speaker": candidate["context_end_speaker"],
                "context_end_dialogue": candidate["context_end_dialogue"],
                "actual_next_excel_row": candidate["actual_next_excel_row"],
                "actual_next_timestamp": candidate["actual_next_timestamp"],
                "actual_next_speaker": candidate["actual_next_speaker"],
                "actual_next_dialogue": candidate["actual_next_dialogue"],
                "candidate_type": candidate["candidate_type"],
                "boundary_type": candidate["boundary_type"],
            }
        )
    return targets


def prediction_prompt(sample_id: str, targets: Sequence[Dict[str, Any]], sanitize_context: bool = False) -> str:
    group_mode = infer_group_mode(sample_id=sample_id)
    session_phrase = "family-therapy transcripts" if group_mode == "family" else "couples-therapy transcripts"
    prompt_targets = []
    for target in targets:
        prior_context = []
        for item in target["prior_context"]:
            prior_context.append(
                {
                    "excel_row": item["excel_row"],
                    "speaker": item["speaker"],
                    "timestamp": item["timestamp"],
                    "dialogue": mask_sensitive_terms(item["dialogue"]) if sanitize_context else item["dialogue"],
                }
            )
        prompt_targets.append(
            {
                "candidate_id": target["candidate_id"],
                "episode_id": target["episode_id"],
                "segment_id": target["segment_id"],
                "decision_after_excel_row": target["context_end_excel_row"],
                "decision_after_timestamp": target["context_end_timestamp"],
                "prior_context": prior_context,
            }
        )
    return (
        f"You are evaluating therapist intervention timing in {session_phrase}.\n"
        "For each target decision point, you only see the transcript context BEFORE that moment.\n"
        "Decide whether the therapist should speak NOW.\n"
        "Return 'yes' if speaking now is the better choice. Return 'no' if the therapist should keep listening.\n"
        "Judge only from the prior context. Do not assume access to the unseen next turn.\n"
        "Use a concise clinical standard: choose 'yes' when an intervention is reasonably timely and helpful now; choose 'no' when continued listening is better.\n"
        "Output JSON only with this schema:\n"
        "{"
        "\"sample_id\":\"detail_1_1\","
        "\"predictions\":["
        "{"
        "\"candidate_id\":\"detail_1_1_pos_13\","
        "\"should_speak\":\"yes\","
        "\"confidence\":0.74,"
        "\"reason\":\"short English reason based only on the prior context\""
        "}"
        "]"
        "}\n"
        f"Sample: {sample_id}\n"
        f"Targets: {json.dumps(prompt_targets, ensure_ascii=False)}"
    )


def call_json_model(
    api_base: str,
    api_key: str,
    model: str,
    prompt: str,
    max_tokens: int,
    disable_thinking: bool = False,
    retries: int = 5,
) -> Dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            payload = {
                "model": model,
                "messages": [{"role": "user", "content": prompt}],
                "temperature": 0,
                "max_tokens": max_tokens,
                "response_format": {"type": "json_object"},
            }
            if disable_thinking:
                payload["enable_thinking"] = False
            response = requests.post(
                f"{api_base}/chat/completions",
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=payload,
                timeout=240,
            )
            if response.status_code == 429:
                raise requests.HTTPError("429", response=response)
            response.raise_for_status()
            content = response.json()["choices"][0]["message"]["content"]
            return json.loads(content)
        except Exception as exc:
            if isinstance(exc, requests.HTTPError) and exc.response is not None:
                response_text = exc.response.text[:2000]
                exc = requests.HTTPError(f"{exc} | response={response_text}", response=exc.response)
            last_error = exc
            if attempt == retries:
                raise exc
            time.sleep(4 * attempt)
    raise RuntimeError(str(last_error))


def request_predictions(
    *,
    api_base: str,
    api_key: str,
    model: str,
    sample_id: str,
    targets: Sequence[Dict[str, Any]],
    max_tokens: int,
    disable_thinking: bool,
) -> Dict[str, Any]:
    try:
        return call_json_model(
            api_base=api_base,
            api_key=api_key,
            model=model,
            prompt=prediction_prompt(sample_id, targets),
            max_tokens=max_tokens,
            disable_thinking=disable_thinking,
        )
    except requests.HTTPError as exc:
        if "data_inspection_failed" not in str(exc):
            raise
        print(f"[retry-sanitized] {sample_id} candidates={len(targets)}", flush=True)
        return call_json_model(
            api_base=api_base,
            api_key=api_key,
            model=model,
            prompt=prediction_prompt(sample_id, targets, sanitize_context=True),
            max_tokens=max_tokens,
            disable_thinking=disable_thinking,
        )


def iter_prediction_items(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    raw_predictions = payload.get("predictions", [])
    if isinstance(raw_predictions, dict):
        raw_predictions = [raw_predictions]
    if not isinstance(raw_predictions, list):
        return []

    normalized: List[Dict[str, Any]] = []
    for item in raw_predictions:
        if isinstance(item, dict):
            normalized.append(item)
            continue
        if isinstance(item, str):
            try:
                parsed = json.loads(item)
            except Exception:
                continue
            if isinstance(parsed, dict):
                normalized.append(parsed)
    return normalized


def normalize_yes_no(value: Any) -> str:
    if isinstance(value, bool):
        return "yes" if value else "no"
    text = clean_text(value).lower()
    if text in {"yes", "y", "true", "1", "speak", "should_speak"}:
        return "yes"
    if text in {"no", "n", "false", "0", "listen", "keep_listening", "do_not_speak"}:
        return "no"
    if "yes" in text:
        return "yes"
    if "no" in text:
        return "no"
    return ""


def parse_confidence(value: Any) -> float | str:
    if value in (None, ""):
        return ""
    try:
        score = float(value)
    except Exception:
        return clean_text(value)
    return max(0.0, min(1.0, score))


def classification_metrics(rows: Sequence[Dict[str, Any]]) -> Dict[str, float]:
    count = len(rows)
    if count == 0:
        return {
            "candidate_count": 0,
            "accuracy": 0.0,
            "precision_yes": 0.0,
            "recall_yes": 0.0,
            "f1_yes": 0.0,
            "precision_no": 0.0,
            "recall_no": 0.0,
            "f1_no": 0.0,
            "balanced_accuracy": 0.0,
            "tp": 0,
            "tn": 0,
            "fp": 0,
            "fn": 0,
        }

    tp = sum(1 for row in rows if row["gold_should_speak"] == "yes" and row["pred_should_speak"] == "yes")
    tn = sum(1 for row in rows if row["gold_should_speak"] == "no" and row["pred_should_speak"] == "no")
    fp = sum(1 for row in rows if row["gold_should_speak"] == "no" and row["pred_should_speak"] == "yes")
    fn = sum(1 for row in rows if row["gold_should_speak"] == "yes" and row["pred_should_speak"] == "no")

    def safe_div(numerator: float, denominator: float) -> float:
        return numerator / denominator if denominator else 0.0

    precision_yes = safe_div(tp, tp + fp)
    recall_yes = safe_div(tp, tp + fn)
    f1_yes = safe_div(2 * precision_yes * recall_yes, precision_yes + recall_yes)
    precision_no = safe_div(tn, tn + fn)
    recall_no = safe_div(tn, tn + fp)
    f1_no = safe_div(2 * precision_no * recall_no, precision_no + recall_no)
    accuracy = safe_div(tp + tn, count)
    balanced_accuracy = (recall_yes + recall_no) / 2 if count else 0.0

    return {
        "candidate_count": count,
        "accuracy": accuracy,
        "precision_yes": precision_yes,
        "recall_yes": recall_yes,
        "f1_yes": f1_yes,
        "precision_no": precision_no,
        "recall_no": recall_no,
        "f1_no": f1_no,
        "balanced_accuracy": balanced_accuracy,
        "tp": tp,
        "tn": tn,
        "fp": fp,
        "fn": fn,
    }


def write_result_workbook(
    output_path: Path,
    overall_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    detail_rows: Sequence[Dict[str, Any]],
    confusion_rows: Sequence[Dict[str, Any]],
) -> None:
    workbook = Workbook()
    ws_overall = workbook.active
    ws_overall.title = "overall"
    ws_overall.append(
        [
            "scope",
            "candidate_count",
            "accuracy",
            "precision_yes",
            "recall_yes",
            "f1_yes",
            "precision_no",
            "recall_no",
            "f1_no",
            "balanced_accuracy",
            "tp",
            "tn",
            "fp",
            "fn",
        ]
    )
    for row in overall_rows:
        ws_overall.append(
            [
                row["scope"],
                row["candidate_count"],
                row["accuracy"],
                row["precision_yes"],
                row["recall_yes"],
                row["f1_yes"],
                row["precision_no"],
                row["recall_no"],
                row["f1_no"],
                row["balanced_accuracy"],
                row["tp"],
                row["tn"],
                row["fp"],
                row["fn"],
            ]
        )

    ws_samples = workbook.create_sheet("samples")
    ws_samples.append(
        [
            "sheet_name",
            "sample_id",
            "candidate_count",
            "accuracy",
            "precision_yes",
            "recall_yes",
            "f1_yes",
            "precision_no",
            "recall_no",
            "f1_no",
            "balanced_accuracy",
            "tp",
            "tn",
            "fp",
            "fn",
        ]
    )
    for row in sample_rows:
        ws_samples.append(
            [
                row["sheet_name"],
                row["sample_id"],
                row["candidate_count"],
                row["accuracy"],
                row["precision_yes"],
                row["recall_yes"],
                row["f1_yes"],
                row["precision_no"],
                row["recall_no"],
                row["f1_no"],
                row["balanced_accuracy"],
                row["tp"],
                row["tn"],
                row["fp"],
                row["fn"],
            ]
        )

    ws_details = workbook.create_sheet("details")
    ws_details.append(
        [
            "sheet_name",
            "sample_id",
            "candidate_id",
            "candidate_type",
            "boundary_type",
            "context_end_excel_row",
            "context_end_timestamp",
            "gold_should_speak",
            "pred_should_speak",
            "correct",
            "confidence",
            "reason",
            "context_end_speaker",
            "context_end_dialogue",
            "actual_next_excel_row",
            "actual_next_timestamp",
            "actual_next_speaker",
            "actual_next_dialogue",
        ]
    )
    for row in detail_rows:
        ws_details.append(
            [
                row["sheet_name"],
                row["sample_id"],
                row["candidate_id"],
                row["candidate_type"],
                row["boundary_type"],
                row["context_end_excel_row"],
                row["context_end_timestamp"],
                row["gold_should_speak"],
                row["pred_should_speak"],
                row["correct"],
                row["confidence"],
                row["reason"],
                row["context_end_speaker"],
                row["context_end_dialogue"],
                row["actual_next_excel_row"],
                row["actual_next_timestamp"],
                row["actual_next_speaker"],
                row["actual_next_dialogue"],
            ]
        )

    ws_conf = workbook.create_sheet("confusions")
    ws_conf.append(["gold_should_speak", "pred_should_speak", "count"])
    for row in confusion_rows:
        ws_conf.append([row["gold_should_speak"], row["pred_should_speak"], row["count"]])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_result_payload(
    model: str,
    api_base: str,
    gold_path: Path,
    detail_rows: Sequence[Dict[str, Any]],
    sample_rows: Sequence[Dict[str, Any]],
    overall_rows: Sequence[Dict[str, Any]],
    confusion_rows: Sequence[Dict[str, Any]],
) -> Dict[str, Any]:
    return {
        "task": "evaluate_therapist_intervention_timing_from_prior_context_only",
        "prediction_model": model,
        "prediction_api_base": api_base,
        "gold_json": str(gold_path),
        "overall": list(overall_rows),
        "samples": list(sample_rows),
        "details": list(detail_rows),
        "confusions": list(confusion_rows),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate a model on therapist intervention timing from prior context only.")
    parser.add_argument("--gold-json", required=True)
    parser.add_argument("--output-json", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--batch-size", type=int, default=4)
    parser.add_argument("--limit-samples", type=int, default=0)
    parser.add_argument("--sample-ids-file", default="", help="Optional path to a text/json file listing sample_ids to include")
    parser.add_argument(
        "--api-base",
        default=os.getenv("OPENAI_API_BASE")
        or os.getenv("OPENAI_BASE_URL")
        or os.getenv("BAILIAN_API_BASE")
        or DEFAULT_API_BASE,
    )
    parser.add_argument(
        "--api-key",
        default=os.getenv("OPENAI_API_KEY") or os.getenv("BAILIAN_API_KEY") or os.getenv("DASHSCOPE_API_KEY"),
    )
    parser.add_argument("--model", default=os.getenv("EVAL_MODEL") or os.getenv("OPENAI_MODEL") or os.getenv("MODEL_NAME") or "")
    parser.add_argument("--max-tokens", type=int, default=2000)
    parser.add_argument("--disable-thinking", action="store_true")
    args = parser.parse_args()

    if not args.api_key:
        raise ValueError("Missing API key.")
    if not args.model:
        raise ValueError("Missing model.")

    gold_path = Path(args.gold_json).expanduser().resolve()
    output_json = Path(args.output_json).expanduser().resolve()
    output_xlsx = Path(args.output_xlsx).expanduser().resolve()
    api_base = normalize_api_base(args.api_base)
    allowed_sample_ids = load_sample_id_filter(args.sample_ids_file)

    gold_samples = load_gold_dataset(gold_path)
    items = sorted(gold_samples.items())
    if args.limit_samples > 0:
        items = items[: args.limit_samples]

    existing_payload: Dict[str, Any] | None = None
    detail_rows: List[Dict[str, Any]] = []
    sample_rows: List[Dict[str, Any]] = []
    if output_json.exists():
        try:
            existing_payload = json.loads(output_json.read_text(encoding="utf-8"))
            detail_rows = list(existing_payload.get("details", []))
            sample_rows = list(existing_payload.get("samples", []))
        except Exception:
            existing_payload = None
            detail_rows = []
            sample_rows = []

    completed_sample_ids = {row["sample_id"] for row in sample_rows}
    by_sample_existing: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        by_sample_existing[row["sample_id"]].append(row)

    for sample_id, sample in items:
        if allowed_sample_ids and sample_id not in allowed_sample_ids:
            continue
        if sample_id in completed_sample_ids:
            print(f"[skip] {sample_id}", flush=True)
            continue

        targets = build_targets(sample)
        if not targets:
            continue

        pred_by_id: Dict[str, Dict[str, Any]] = {}
        for batch in batched(targets, args.batch_size):
            payload = request_predictions(
                api_base=api_base,
                api_key=args.api_key,
                model=args.model,
                sample_id=sample_id,
                targets=batch,
                max_tokens=args.max_tokens,
                disable_thinking=args.disable_thinking,
            )
            for item in iter_prediction_items(payload):
                candidate_id = clean_text(item.get("candidate_id"))
                if candidate_id:
                    pred_by_id[candidate_id] = item
            time.sleep(1)

        missing_targets = [item for item in targets if item["candidate_id"] not in pred_by_id]
        for target in missing_targets:
            payload = request_predictions(
                api_base=api_base,
                api_key=args.api_key,
                model=args.model,
                sample_id=sample_id,
                targets=[target],
                max_tokens=max(700, min(args.max_tokens, 1000)),
                disable_thinking=args.disable_thinking,
            )
            for item in iter_prediction_items(payload):
                candidate_id = clean_text(item.get("candidate_id"))
                if candidate_id:
                    pred_by_id[candidate_id] = item
            time.sleep(1)

        sample_detail_rows: List[Dict[str, Any]] = []
        for target in targets:
            pred = pred_by_id.get(target["candidate_id"], {})
            pred_should_speak = normalize_yes_no(pred.get("should_speak"))
            if pred_should_speak not in {"yes", "no"}:
                pred_should_speak = "no"
            sample_detail_rows.append(
                {
                    "sheet_name": target["sheet_name"],
                    "sample_id": sample_id,
                    "candidate_id": target["candidate_id"],
                    "candidate_type": target["candidate_type"],
                    "boundary_type": target["boundary_type"],
                    "context_end_excel_row": target["context_end_excel_row"],
                    "context_end_timestamp": target["context_end_timestamp"],
                    "gold_should_speak": target["gold_should_speak"],
                    "pred_should_speak": pred_should_speak,
                    "correct": 1.0 if pred_should_speak == target["gold_should_speak"] else 0.0,
                    "confidence": parse_confidence(pred.get("confidence")),
                    "reason": clean_text(pred.get("reason")),
                    "context_end_speaker": target["context_end_speaker"],
                    "context_end_dialogue": target["context_end_dialogue"],
                    "actual_next_excel_row": target["actual_next_excel_row"],
                    "actual_next_timestamp": target["actual_next_timestamp"],
                    "actual_next_speaker": target["actual_next_speaker"],
                    "actual_next_dialogue": target["actual_next_dialogue"],
                }
            )

        detail_rows.extend(sample_detail_rows)
        sample_metrics = classification_metrics(sample_detail_rows)
        sample_row = {
            "sheet_name": sample["sheet_name"],
            "sample_id": sample_id,
            **sample_metrics,
        }
        sample_rows.append(sample_row)
        completed_sample_ids.add(sample_id)

        overall_rows = [{"scope": "overall", **classification_metrics(detail_rows)}]
        confusion_counter = Counter((row["gold_should_speak"], row["pred_should_speak"]) for row in detail_rows)
        confusion_rows = [
            {"gold_should_speak": gold, "pred_should_speak": pred, "count": count}
            for (gold, pred), count in sorted(confusion_counter.items())
        ]
        payload = build_result_payload(
            model=args.model,
            api_base=api_base,
            gold_path=gold_path,
            detail_rows=detail_rows,
            sample_rows=sample_rows,
            overall_rows=overall_rows,
            confusion_rows=confusion_rows,
        )
        output_json.parent.mkdir(parents=True, exist_ok=True)
        output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(
            f"[done] {sample_id}: candidates={sample_metrics['candidate_count']} "
            f"acc={sample_metrics['accuracy']:.3f} f1_yes={sample_metrics['f1_yes']:.3f}",
            flush=True,
        )

    overall_rows = [{"scope": "overall", **classification_metrics(detail_rows)}]
    confusion_counter = Counter((row["gold_should_speak"], row["pred_should_speak"]) for row in detail_rows)
    confusion_rows = [
        {"gold_should_speak": gold, "pred_should_speak": pred, "count": count}
        for (gold, pred), count in sorted(confusion_counter.items())
    ]
    payload = build_result_payload(
        model=args.model,
        api_base=api_base,
        gold_path=gold_path,
        detail_rows=detail_rows,
        sample_rows=sample_rows,
        overall_rows=overall_rows,
        confusion_rows=confusion_rows,
    )
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_result_workbook(output_xlsx, overall_rows, sample_rows, detail_rows, confusion_rows)
    print(f"[written] {output_json}")
    print(f"[written] {output_xlsx}")


if __name__ == "__main__":
    main()
